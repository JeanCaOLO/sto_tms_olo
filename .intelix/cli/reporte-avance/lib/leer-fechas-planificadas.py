"""Extrae las fechas planificadas de la estimación FINAL en formato JSON."""

import datetime as dt
import json
import sys

import openpyxl


def fecha_iso(valor):
    if isinstance(valor, (dt.datetime, dt.date)):
        return valor.strftime("%Y-%m-%d")
    texto = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(texto, formato).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def main(ruta):
    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    try:
        hoja = libro["Plan"]
        cabecera = next(hoja.iter_rows(values_only=True))
        for fila in hoja.iter_rows(values_only=True):
            normalizada = [str(valor or "").strip().lower() for valor in fila]
            if "id tracker" in normalizada and "fin" in normalizada:
                cabecera = fila
                break

        columnas = {str(valor or "").strip().lower(): indice for indice, valor in enumerate(cabecera)}
        id_tracker = columnas["id tracker"]
        fin = columnas["fin"]
        fechas = {}
        for fila in hoja.iter_rows(min_row=5, values_only=True):
            identificador = str(fila[id_tracker] or "").strip()
            if identificador.upper().startswith("VER-"):
                fecha = fecha_iso(fila[fin])
                if fecha:
                    fechas[identificador] = fecha
        print(json.dumps(fechas))
    finally:
        libro.close()


if __name__ == "__main__":
    main(sys.argv[1])
