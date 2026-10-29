#!/usr/bin/env python3
"""Control de iteraciones de Inception + impacto en Construction (con HH).

Fuente de verdad: `.aidlc/audit/inception-iterations.jsonl` (una linea JSON por evento).
Reportes: `.aidlc/audit/inception-report.md` y `.html`.
Contadores en `.aidlc/aidlc-state.md` via `sync-state`.

Comandos:
  add | list | status | report | sync-state | simulate
"""

from __future__ import annotations

import argparse
import html
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
AUDIT = ROOT / ".aidlc" / "audit"
LEDGER = AUDIT / "inception-iterations.jsonl"
STATE = ROOT / ".aidlc" / "aidlc-state.md"
REPORT_MD = AUDIT / "inception-report.md"
REPORT_HTML = AUDIT / "inception-report.html"
EXAMPLES = AUDIT / "examples" / "simulacion-inception"
HH_CATALOG_CANDIDATES = (
    AUDIT / "hh-catalog.json",
    ROOT / ".intelix" / "outputs" / "hh-catalog.json",
)

PHASES = ("inception", "construction_impact")
KINDS = (
    "gate_rejected",
    "po_scope_change",
    "tech_discovery",
    "reclassification",
    "other",
)


def today_iso() -> str:
    return date.today().isoformat()


def utc_now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def ensure_audit() -> None:
    AUDIT.mkdir(parents=True, exist_ok=True)


def load_ledger(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    rows: list[dict[str, Any]] = []
    for i, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        try:
            rows.append(json.loads(line))
        except json.JSONDecodeError as exc:
            raise SystemExit(f"Ledger invalido en {path}:{i}: {exc}") from exc
    return rows


def append_event(path: Path, event: dict[str, Any]) -> None:
    ensure_audit()
    with path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(event, ensure_ascii=False) + "\n")


def split_csv(value: str | None) -> list[str]:
    if not value:
        return []
    return [p.strip() for p in value.split(",") if p.strip()]


def next_id(rows: list[dict[str, Any]], phase: str) -> str:
    prefix = "INC-ITER" if phase == "inception" else "SCOPE-RW"
    nums = []
    for row in rows:
        m = re.match(rf"^{re.escape(prefix)}-(\d+)$", str(row.get("id", "")))
        if m:
            nums.append(int(m.group(1)))
    return f"{prefix}-{(max(nums) if nums else 0) + 1:02d}"


def find_plan_xlsx(explicit: str | None = None) -> Path | None:
    if explicit:
        p = Path(explicit)
        return p if p.is_file() else None
    patterns = (
        "PLAN-ESTIMACION-*-FINAL.xlsx",
        "**/PLAN-ESTIMACION-*-FINAL.xlsx",
        ".intelix/outputs/**/PLAN-ESTIMACION-*-FINAL.xlsx",
    )
    for pat in patterns:
        hits = sorted(ROOT.glob(pat), key=lambda x: x.stat().st_mtime, reverse=True)
        if hits:
            return hits[0]
    return None


def load_hh_catalog(path: Path | None = None, plan: Path | None = None) -> dict[str, float]:
    """Mapa ID tarea → HH planificadas.

    Orden: --hh-catalog JSON → openpyxl sobre PLAN-ESTIMACION → vacio.
    JSON: {\"U1-T2\": 4.5, ...} o {\"tasks\": {\"U1-T2\": {\"hh\": 4.5}}}
    """
    catalog: dict[str, float] = {}
    candidates = [path] if path else list(HH_CATALOG_CANDIDATES)
    for cand in candidates:
        if cand and cand.is_file():
            data = json.loads(cand.read_text(encoding="utf-8"))
            if isinstance(data, dict) and "tasks" in data:
                for tid, meta in (data.get("tasks") or {}).items():
                    if isinstance(meta, dict) and "hh" in meta:
                        catalog[str(tid)] = float(meta["hh"])
                    elif isinstance(meta, (int, float)):
                        catalog[str(tid)] = float(meta)
            elif isinstance(data, dict):
                for tid, val in data.items():
                    if isinstance(val, (int, float)):
                        catalog[str(tid)] = float(val)
            if catalog:
                return catalog

    xlsx = plan or find_plan_xlsx()
    if not xlsx:
        return catalog
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return catalog

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if "Plan" not in wb.sheetnames:
        return catalog
    ws = wb["Plan"]
    # Cabecera tipica fila 4: buscar columnas Duracion e ID Tracker / Id
    header_row = 4
    headers = {
        str(ws.cell(header_row, c).value or "").strip().lower(): c
        for c in range(1, (ws.max_column or 1) + 1)
    }
    col_dur = None
    for key in ("duracion (horas)", "duracion", "duración (horas)", "duración"):
        if key in headers:
            col_dur = headers[key]
            break
    col_tid = None
    for key in ("id tracker", "id jerarquico", "id jerárquico", "id"):
        if key in headers:
            col_tid = headers[key]
            break
    if not col_dur or not col_tid:
        return catalog
    for r in range(header_row + 1, (ws.max_row or header_row) + 1):
        tid = ws.cell(r, col_tid).value
        dur = ws.cell(r, col_dur).value
        if tid is None or dur is None:
            continue
        try:
            catalog[str(tid).strip()] = float(dur)
        except (TypeError, ValueError):
            continue
    return catalog


def compute_hh_block(
    tasks_reopened: list[str],
    tasks_added: list[str],
    tasks_skipped: list[str],
    hh_delta: float | None,
    catalog: dict[str, float],
    source_hint: str = "",
) -> dict[str, Any]:
    by_task: dict[str, float] = {}
    for tid in tasks_reopened + tasks_added:
        if tid in catalog:
            by_task[tid] = catalog[tid]
    skipped_hh = 0.0
    for tid in tasks_skipped:
        if tid in catalog:
            skipped_hh += catalog[tid]
            by_task[f"-{tid}"] = -catalog[tid]
    derived = sum(by_task.values())
    if hh_delta is not None:
        estimated = float(hh_delta)
        source = source_hint or "manual"
    elif by_task:
        estimated = round(derived, 2)
        source = source_hint or "catalog/plan"
    else:
        estimated = 0.0
        source = source_hint or "none"
    return {
        "estimated": round(estimated, 2),
        "source": source,
        "by_task": by_task,
        "skipped_hh": round(skipped_hh, 2),
    }


def event_hh(event: dict[str, Any]) -> float:
    hh = event.get("hh") or {}
    try:
        return float(hh.get("estimated") or 0)
    except (TypeError, ValueError):
        return 0.0


def build_event(
    args: argparse.Namespace,
    rows: list[dict[str, Any]],
    catalog: dict[str, float] | None = None,
) -> dict[str, Any]:
    phase = args.phase
    eid = args.id or next_id(rows, phase)
    tasks_reopened = split_csv(args.tasks_reopened)
    tasks_added = split_csv(args.tasks_added)
    tasks_skipped = split_csv(args.tasks_skipped)
    impact = {
        "affected": bool(args.affects_construction)
        or phase == "construction_impact"
        or bool(args.developers)
        or bool(args.branches)
        or bool(tasks_reopened)
        or bool(args.develop_effect),
        "developers": split_csv(args.developers),
        "branches": split_csv(args.branches),
        "develop_effect": args.develop_effect or "",
        "tasks_reopened": tasks_reopened,
        "tasks_added": tasks_added,
        "tasks_skipped": tasks_skipped,
    }
    cat = catalog if catalog is not None else {}
    hh_delta = float(args.hh_delta) if getattr(args, "hh_delta", None) not in (None, "") else None
    hh = compute_hh_block(
        tasks_reopened,
        tasks_added,
        tasks_skipped,
        hh_delta,
        cat,
        source_hint="manual" if hh_delta is not None else "",
    )
    return {
        "id": eid,
        "phase": phase,
        "kind": args.kind,
        "opened_at": args.opened_at or today_iso(),
        "closed_at": args.closed_at or "",
        "requested_by": args.requested_by or "",
        "approved_by": args.approved_by or "",
        "summary": args.summary,
        "artifacts_changed": split_csv(args.artifacts),
        "construction_impact": impact,
        "hh": hh,
        "evidence": {
            "issue": args.issue or "",
            "mr": args.mr or "",
        },
        "recorded_at": utc_now(),
    }


def cmd_add(args: argparse.Namespace) -> int:
    ensure_audit()
    rows = load_ledger(LEDGER)
    catalog = load_hh_catalog(
        Path(args.hh_catalog) if args.hh_catalog else None,
        find_plan_xlsx(args.plan if getattr(args, "plan", "") else None),
    )
    event = build_event(args, rows, catalog)
    if any(r.get("id") == event["id"] for r in rows):
        raise SystemExit(f"Ya existe el id {event['id']} en el ledger.")
    append_event(LEDGER, event)
    print(
        f"Registrado {event['id']} → {LEDGER.relative_to(ROOT)} "
        f"(HH estimadas: {event['hh']['estimated']})"
    )
    if args.sync_state:
        cmd_sync_state(argparse.Namespace(ledger=str(LEDGER)))
    return 0


def cmd_list(args: argparse.Namespace) -> int:
    ledger = Path(args.ledger) if args.ledger else LEDGER
    rows = load_ledger(ledger)
    if args.phase:
        rows = [r for r in rows if r.get("phase") == args.phase]
    if not rows:
        print("Ledger vacio.")
        return 0
    print(f"{'ID':<14} {'Fase':<20} {'Apertura':<12} {'Cierre':<12} {'HH':>7}  Resumen")
    for r in sorted(rows, key=lambda x: x.get("opened_at") or ""):
        print(
            f"{str(r.get('id')):<14} {str(r.get('phase')):<20} "
            f"{str(r.get('opened_at') or '—'):<12} {str(r.get('closed_at') or '—'):<12} "
            f"{event_hh(r):7.2f}  {str(r.get('summary') or '')[:60]}"
        )
    return 0


def cmd_status(args: argparse.Namespace) -> int:
    ledger = Path(args.ledger) if args.ledger else LEDGER
    stats = summarize(load_ledger(ledger))
    print(f"Proyecto: {project_name()}")
    print(f"Iteraciones Inception:     {stats['inception_iterations']}")
    print(f"Impactos Construction:     {stats['construction_impacts']}")
    print(f"HH impacto Inception:      {stats['hh_inception']:.2f}")
    print(f"HH impacto Construction:   {stats['hh_construction']:.2f}")
    print(f"HH impacto total:          {stats['hh_total']:.2f}")
    print(f"Primera fecha:             {stats['first_at'] or '—'}")
    print(f"Ultima fecha:              {stats['last_at'] or '—'}")
    return 0


def summarize(rows: list[dict[str, Any]]) -> dict[str, Any]:
    inception = [r for r in rows if r.get("phase") == "inception"]
    impacts = [
        r
        for r in rows
        if r.get("phase") == "construction_impact"
        or (r.get("construction_impact") or {}).get("affected")
    ]
    dates = []
    for r in rows:
        for key in ("opened_at", "closed_at"):
            val = r.get(key) or ""
            if val:
                dates.append(val)
    by_dev: dict[str, list[str]] = defaultdict(list)
    by_branch: dict[str, list[str]] = defaultdict(list)
    develop_notes: list[tuple[str, str]] = []
    hh_by_dev: dict[str, float] = defaultdict(float)
    kind_counts: dict[str, int] = defaultdict(int)
    kind_hh: dict[str, float] = defaultdict(float)
    for r in rows:
        kind = str(r.get("kind") or "other")
        kind_counts[kind] += 1
        kind_hh[kind] += event_hh(r)
    for r in impacts:
        ci = r.get("construction_impact") or {}
        hh = event_hh(r)
        devs = ci.get("developers") or []
        for d in devs:
            by_dev[d].append(r["id"])
            if devs:
                hh_by_dev[d] += hh / len(devs)
        for b in ci.get("branches") or []:
            by_branch[b].append(r["id"])
        effect = (ci.get("develop_effect") or "").strip()
        if effect:
            develop_notes.append((r["id"], effect))
    hh_inception = sum(event_hh(r) for r in inception)
    hh_construction = sum(event_hh(r) for r in impacts)
    return {
        "total": len(rows),
        "inception_iterations": len(inception),
        "construction_impacts": len(impacts),
        "hh_inception": round(hh_inception, 2),
        "hh_construction": round(hh_construction, 2),
        "hh_total": round(sum(event_hh(r) for r in rows), 2),
        "first_at": min(dates) if dates else "",
        "last_at": max(dates) if dates else "",
        "by_dev": dict(by_dev),
        "hh_by_dev": {k: round(v, 2) for k, v in hh_by_dev.items()},
        "by_branch": dict(by_branch),
        "develop_notes": develop_notes,
        "kind_counts": dict(kind_counts),
        "kind_hh": {k: round(v, 2) for k, v in kind_hh.items()},
        "rows": rows,
        "inception_rows": inception,
        "impact_rows": impacts,
    }


KIND_LABELS = {
    "gate_rejected": "Gate Inception rechazado",
    "po_scope_change": "Cambio de alcance (PO/negocio)",
    "tech_discovery": "Hallazgo tecnico",
    "reclassification": "Reclasificacion Menor/Intermedio/Mayor",
    "other": "Otro",
}

HIGHLIGHTS = [
    (
        "Avance oficial = develop",
        "Igual que `intelix-consolidado-avance`: el % oficial solo cuenta lo integrado en `develop`. "
        "Este reporte explica el rework que obligo a reabrir features antes de integrar.",
    ),
    (
        "Discovery + Gate 0",
        "La entrada a Inception sigue el contrato `Product-Definition/` y el sign-off en "
        "`.aidlc/aidlc-state.md`. Las vueltas de Inception se cuentan *despues* de Gate 0.",
    ),
    (
        "Rigor proporcional",
        "Clasificacion Menor / Intermedio / Mayor define gates tipicos. Un rework puede "
        "forzar reclasificacion explicita (`kind=reclassification`), nunca en silencio.",
    ),
    (
        "REWORK_FACTOR vs rework real",
        "`intelix-estimacion` reserva ~5% de margen planificado. El ledger mide HH de "
        "rework **real** por evento para contrastar plan vs ejecucion.",
    ),
    (
        "Seguridad y estandares",
        "Vueltas por STRIDE / baseline R1–R14 o por incumplimiento de `docs/estandares/` "
        "se registran como `gate_rejected` o `tech_discovery` con evidencia en issue/MR.",
    ),
    (
        "Multi-dev / GitLab",
        "Impactos listan desarrolladores, ramas `feature/*` y efecto en `develop`. "
        "Issue template: *Inception Iteration or Scope Change*.",
    ),
]


def read_state_field(text: str, key: str) -> str:
    m = re.search(rf"- \*\*{re.escape(key)}\*\*: `([^`]*)`", text)
    return (m.group(1).strip() if m else "") or ""


def load_project_context(overrides: dict[str, str] | None = None) -> dict[str, str]:
    ctx = {
        "name": ROOT.name,
        "entry_mode": "",
        "classification": "",
        "gate_status": "",
        "inception_gate_status": "",
        "aidlc_version": "",
        "has_product_definition": "si" if (ROOT / "Product-Definition").is_dir() else "no",
        "has_gate0_script": "si"
        if (ROOT / ".intelix" / "scripts" / "validate_discovery_gate.py").is_file()
        else "no",
        "has_verificacion": "si" if (ROOT / "scripts" / "verificacion").is_dir() else "no",
        "has_consolidado_skill": "si"
        if any(ROOT.glob("**/intelix-consolidado-avance/SKILL.md"))
        or (ROOT / "templates" / "skills" / "intelix-consolidado-avance" / "SKILL.md").is_file()
        else "no",
    }
    # Desde el arbol de plantilla (este script vive en templates/.intelix/scripts)
    if (ROOT / "skills" / "intelix-consolidado-avance" / "SKILL.md").is_file():
        ctx["has_consolidado_skill"] = "si"
    if STATE.is_file():
        text = STATE.read_text(encoding="utf-8")
        name = read_state_field(text, "name")
        if name and name != "{{projectName}}":
            ctx["name"] = name
        for key in (
            "entry_mode",
            "classification",
            "gate_status",
            "inception_gate_status",
        ):
            ctx[key] = read_state_field(text, key)
    config_path = ROOT / ".intelix" / "config.json"
    if config_path.is_file():
        try:
            config = json.loads(config_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            config = {}
        version = config.get("aidlcVersion") or (config.get("aidlc") or {}).get("version")
        if version:
            ctx["aidlc_version"] = str(version)
    if overrides:
        ctx.update({k: v for k, v in overrides.items() if v is not None})
    return ctx


def fmt_hh(v: float) -> str:
    return f"{v:.2f}" if v else "—"


def render_md(
    project: str,
    stats: dict[str, Any],
    generated: str,
    ctx: dict[str, str] | None = None,
) -> str:
    ctx = ctx or load_project_context({"name": project})
    lines = [
        "# Reporte de gestion — Iteraciones de Inception",
        "",
        "> Plantilla INTELIX · IA-DLC · control de vueltas post Gate 0",
        "",
        f"**Proyecto:** {project}  ",
        f"**Generado:** {generated}  ",
        f"**Ledger:** `.aidlc/audit/inception-iterations.jsonl`",
        "",
        "## Contexto del proyecto (ADN plantilla)",
        "",
        "| Elemento plantilla | Valor |",
        "|--------------------|-------|",
        f"| Entry mode (Discovery) | `{ctx.get('entry_mode') or '—'}` |",
        f"| Gate 0 (`gate_status`) | `{ctx.get('gate_status') or '—'}` |",
        f"| Clasificacion | `{ctx.get('classification') or '—'}` |",
        f"| Gate Inception | `{ctx.get('inception_gate_status') or '—'}` |",
        f"| Pin AI-DLC | `{ctx.get('aidlc_version') or '—'}` |",
        f"| `Product-Definition/` | {ctx.get('has_product_definition')} |",
        f"| Script Gate 0 | {ctx.get('has_gate0_script')} |",
        f"| Motor certificacion | {ctx.get('has_verificacion')} |",
        "",
        "## Elementos destacados de la plantilla",
        "",
    ]
    for title, body in HIGHLIGHTS:
        lines.append(f"- **{title}.** {body}")
    lines += [
        "",
        "## KPI",
        "",
        "| Indicador | Valor |",
        "|-----------|-------|",
        f"| Iteraciones en Inception | **{stats['inception_iterations']}** |",
        f"| Eventos con impacto en Construction / develop | **{stats['construction_impacts']}** |",
        f"| HH impacto Inception | **{fmt_hh(stats['hh_inception'])}** |",
        f"| HH impacto Construction | **{fmt_hh(stats['hh_construction'])}** |",
        f"| HH impacto total (rework real) | **{fmt_hh(stats['hh_total'])}** |",
        f"| Primera fecha registrada | {stats['first_at'] or '—'} |",
        f"| Ultima fecha registrada | {stats['last_at'] or '—'} |",
        "",
        "## Taxonomia de disparadores (plantilla)",
        "",
        "| Disparador | Eventos | HH |",
        "|------------|---------|----|",
    ]
    if not stats["kind_counts"]:
        lines.append("| — | 0 | — |")
    else:
        for kind, count in sorted(stats["kind_counts"].items(), key=lambda x: (-x[1], x[0])):
            label = KIND_LABELS.get(kind, kind)
            lines.append(
                f"| `{kind}` — {label} | {count} | {fmt_hh(stats['kind_hh'].get(kind, 0))} |"
            )
    lines += [
        "",
        "## Posicion en el ciclo AI-DLC",
        "",
        "```text",
        "Initialization → Ideation → [Gate 0] → Inception* → Construction* → Operation",
        "                                 ↑              ↑",
        "                          Product-Definition   INC-ITER / SCOPE-RW (este reporte)",
        "```",
        "",
        "## Linea de tiempo (Inception)",
        "",
    ]
    if not stats["inception_rows"]:
        lines.append("_Sin iteraciones de Inception registradas._")
        lines.append("")
    else:
        lines += [
            "| ID | Apertura | Cierre | Disparador | HH | Resumen | Evidencia |",
            "|----|----------|--------|------------|----|---------|-----------|",
        ]
        for r in sorted(stats["inception_rows"], key=lambda x: x.get("opened_at") or ""):
            ev = r.get("evidence") or {}
            evid = ", ".join(x for x in (ev.get("issue"), ev.get("mr")) if x) or "—"
            lines.append(
                f"| {r.get('id')} | {r.get('opened_at') or '—'} | {r.get('closed_at') or '—'} | "
                f"{r.get('kind')} | {fmt_hh(event_hh(r))} | "
                f"{r.get('summary', '').replace('|', '/')} | {evid} |"
            )
        lines.append("")

    lines += [
        "## Impacto en Construction y rama `develop`",
        "",
        "Principio plantilla: el avance **oficial** solo mide `develop` "
        "(`intelix-consolidado-avance`). Aqui se ve el costo en HH y tareas "
        "reabiertas/nuevas que ese rework genero en features.",
        "",
    ]
    if not stats["impact_rows"]:
        lines.append("_Ningun evento con impacto en Construction._")
        lines.append("")
    else:
        lines += [
            "| ID | Fecha | Devs | Ramas | HH | Efecto en develop | Reabiertas | Nuevas |",
            "|----|-------|------|-------|----|-------------------|------------|--------|",
        ]
        for r in sorted(stats["impact_rows"], key=lambda x: x.get("opened_at") or ""):
            ci = r.get("construction_impact") or {}
            lines.append(
                "| {id} | {dt} | {devs} | {br} | {hh} | {eff} | {reo} | {add} |".format(
                    id=r.get("id"),
                    dt=r.get("opened_at") or "—",
                    devs=", ".join(ci.get("developers") or []) or "—",
                    br=", ".join(ci.get("branches") or []) or "—",
                    hh=fmt_hh(event_hh(r)),
                    eff=(ci.get("develop_effect") or "—").replace("|", "/"),
                    reo=", ".join(ci.get("tasks_reopened") or []) or "—",
                    add=", ".join(ci.get("tasks_added") or []) or "—",
                )
            )
        lines.append("")

    lines += ["## Por desarrollador (eventos y HH prorrateadas)", ""]
    if not stats["by_dev"]:
        lines.append("_Sin desarrolladores impactados._")
        lines.append("")
    else:
        for dev, ids in sorted(stats["by_dev"].items()):
            hh = stats["hh_by_dev"].get(dev, 0)
            lines.append(f"- **{dev}**: {len(ids)} evento(s), ~{fmt_hh(hh)} HH — {', '.join(ids)}")
        lines.append("")

    lines += ["## Por rama feature", ""]
    if not stats["by_branch"]:
        lines.append("_Sin ramas feature impactadas._")
        lines.append("")
    else:
        for br, ids in sorted(stats["by_branch"].items()):
            lines.append(f"- `{br}`: {', '.join(ids)}")
        lines.append("")

    lines += ["## Efectos acumulados en `develop`", ""]
    if not stats["develop_notes"]:
        lines.append("_Sin notas de efecto en develop._")
    else:
        for eid, note in stats["develop_notes"]:
            lines.append(f"- **{eid}**: {note}")
    lines += [
        "",
        "## Capacidades relacionadas (plantilla)",
        "",
        "| Capacidad | Para que cruzar |",
        "|-----------|-----------------|",
        "| `intelix-estimacion` | Baseline HH y `REWORK_FACTOR` planificado |",
        "| `intelix-task-tracker` | IDs de tareas reabiertas / nuevas / omitidas |",
        "| `intelix-consolidado-avance` | % oficial develop vs consolidado multi-rama |",
        "| `intelix-seguimiento-plan` | Plan vs real en fechas |",
        "| `intelix-certificacion` | Dictamen de estandares antes de QA |",
        "| Gate 0 / `Product-Definition/` | Contrato previo a Inception |",
        "",
        "---",
        "",
        "_Guia: `docs/guides/inception-iterations.md` · "
        "Skill: `intelix-inception-iterations` · "
        "Comandos: `add` / `list` / `status` / `report` / `sync-state`._",
        "",
    ]
    return "\n".join(lines)


def render_html(
    project: str,
    stats: dict[str, Any],
    generated: str,
    ctx: dict[str, str] | None = None,
) -> str:
    ctx = ctx or load_project_context({"name": project})

    def esc(s: Any) -> str:
        return html.escape(str(s if s is not None else ""))

    def rows_inception() -> str:
        if not stats["inception_rows"]:
            return "<tr><td colspan='7'><em>Sin iteraciones</em></td></tr>"
        out = []
        for r in sorted(stats["inception_rows"], key=lambda x: x.get("opened_at") or ""):
            ev = r.get("evidence") or {}
            evid = ", ".join(x for x in (ev.get("issue"), ev.get("mr")) if x) or "—"
            out.append(
                "<tr>"
                f"<td>{esc(r.get('id'))}</td>"
                f"<td>{esc(r.get('opened_at') or '—')}</td>"
                f"<td>{esc(r.get('closed_at') or '—')}</td>"
                f"<td><span class='tag'>{esc(r.get('kind'))}</span></td>"
                f"<td><strong>{esc(fmt_hh(event_hh(r)))}</strong></td>"
                f"<td>{esc(r.get('summary'))}</td>"
                f"<td>{esc(evid)}</td>"
                "</tr>"
            )
        return "\n".join(out)

    def rows_impact() -> str:
        if not stats["impact_rows"]:
            return "<tr><td colspan='8'><em>Sin impactos</em></td></tr>"
        out = []
        for r in sorted(stats["impact_rows"], key=lambda x: x.get("opened_at") or ""):
            ci = r.get("construction_impact") or {}
            out.append(
                "<tr>"
                f"<td>{esc(r.get('id'))}</td>"
                f"<td>{esc(r.get('opened_at') or '—')}</td>"
                f"<td>{esc(', '.join(ci.get('developers') or []) or '—')}</td>"
                f"<td><code>{esc(', '.join(ci.get('branches') or []) or '—')}</code></td>"
                f"<td><strong>{esc(fmt_hh(event_hh(r)))}</strong></td>"
                f"<td>{esc(ci.get('develop_effect') or '—')}</td>"
                f"<td>{esc(', '.join(ci.get('tasks_reopened') or []) or '—')}</td>"
                f"<td>{esc(', '.join(ci.get('tasks_added') or []) or '—')}</td>"
                "</tr>"
            )
        return "\n".join(out)

    def rows_kinds() -> str:
        if not stats["kind_counts"]:
            return "<tr><td colspan='3'><em>Sin eventos</em></td></tr>"
        out = []
        for kind, count in sorted(stats["kind_counts"].items(), key=lambda x: (-x[1], x[0])):
            out.append(
                "<tr>"
                f"<td><code>{esc(kind)}</code><div class='muted'>{esc(KIND_LABELS.get(kind, kind))}</div></td>"
                f"<td>{count}</td>"
                f"<td><strong>{esc(fmt_hh(stats['kind_hh'].get(kind, 0)))}</strong></td>"
                "</tr>"
            )
        return "\n".join(out)

    highlights = "".join(
        f"<div class='hl'><h3>{esc(t)}</h3><p>{esc(b)}</p></div>" for t, b in HIGHLIGHTS
    )
    ctx_rows = "".join(
        f"<tr><td>{esc(label)}</td><td><code>{esc(val or '—')}</code></td></tr>"
        for label, val in (
            ("Entry mode (Discovery)", ctx.get("entry_mode")),
            ("Gate 0", ctx.get("gate_status")),
            ("Clasificacion", ctx.get("classification")),
            ("Gate Inception", ctx.get("inception_gate_status")),
            ("Pin AI-DLC", ctx.get("aidlc_version")),
            ("Product-Definition/", ctx.get("has_product_definition")),
            ("Script Gate 0", ctx.get("has_gate0_script")),
            ("Motor certificacion", ctx.get("has_verificacion")),
        )
    )
    dev_lis = (
        "".join(
            f"<li><strong>{esc(d)}</strong>: {len(ids)} evento(s), "
            f"~{esc(fmt_hh(stats['hh_by_dev'].get(d, 0)))} HH — {esc(', '.join(ids))}</li>"
            for d, ids in sorted(stats["by_dev"].items())
        )
        or "<li><em>Ninguno</em></li>"
    )
    br_lis = (
        "".join(
            f"<li><code>{esc(b)}</code>: {esc(', '.join(ids))}</li>"
            for b, ids in sorted(stats["by_branch"].items())
        )
        or "<li><em>Ninguna</em></li>"
    )
    dev_notes = (
        "".join(f"<li><strong>{esc(i)}</strong>: {esc(n)}</li>" for i, n in stats["develop_notes"])
        or "<li><em>Sin notas</em></li>"
    )

    return f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="utf-8"/>
<title>Iteraciones Inception — {esc(project)}</title>
<style>
  :root {{
    --navy:#0A1F44; --blue:#1E5AE8; --cyan:#00C2D1; --emerald:#12B886;
    --amber:#F5A524; --line:#E2E8F4; --muted:#6B7A99; --bg:#F5F7FC; --card:#fff;
  }}
  body {{ font-family:'Segoe UI',system-ui,sans-serif; margin:0; background:var(--bg); color:var(--navy); }}
  .banner {{
    background:linear-gradient(135deg,#0A1F44 0%,#16305e 55%,#1E5AE8 100%);
    color:#fff; padding:28px 20px 24px;
  }}
  .banner .brand {{
    display:inline-flex; align-items:center; gap:8px; font-size:12px; font-weight:700;
    letter-spacing:.08em; text-transform:uppercase; color:#8fe6ee; margin-bottom:10px;
  }}
  .banner .brand i {{
    width:10px; height:10px; border-radius:3px;
    background:linear-gradient(135deg,var(--cyan),var(--blue)); display:inline-block;
  }}
  .banner h1 {{ margin:0 0 8px; font-size:28px; letter-spacing:-.3px; }}
  .banner .meta {{ color:#b8c7e6; font-size:13.5px; }}
  main {{ max-width:1100px; margin:0 auto; padding:24px 20px 48px; }}
  .kpis {{ display:grid; grid-template-columns:repeat(3,1fr); gap:12px; margin:18px 0 22px; }}
  .kpi {{ background:var(--card); border:1px solid var(--line); border-radius:12px; padding:16px; }}
  .kpi .n {{ font-size:28px; font-weight:800; color:var(--blue); }}
  .kpi .l {{ font-size:12px; color:var(--muted); margin-top:4px; }}
  .hl-grid {{ display:grid; grid-template-columns:1fr 1fr; gap:12px; margin-bottom:16px; }}
  .hl {{ background:var(--card); border:1px solid var(--line); border-left:4px solid var(--cyan);
        border-radius:12px; padding:14px 16px; }}
  .hl h3 {{ margin:0 0 6px; font-size:13.5px; }}
  .hl p {{ margin:0; font-size:12.5px; color:var(--muted); line-height:1.45; }}
  section {{ background:var(--card); border:1px solid var(--line); border-radius:14px;
             padding:18px 20px; margin-bottom:16px; }}
  h2 {{ font-size:17px; margin:0 0 12px; }}
  table {{ width:100%; border-collapse:collapse; font-size:13px; }}
  th {{ text-align:left; background:var(--navy); color:#dfe8ff; padding:10px 12px; }}
  td {{ padding:9px 12px; border-bottom:1px solid var(--line); vertical-align:top; }}
  ul {{ margin:0; padding-left:18px; }}
  code {{ font-size:12px; }}
  .tag {{ display:inline-block; background:#EAF1FF; color:var(--blue); border-radius:6px;
          padding:2px 8px; font-size:11px; font-weight:700; }}
  .muted {{ color:var(--muted); font-size:11.5px; margin-top:2px; }}
  .cycle {{
    font-family:ui-monospace,Consolas,monospace; font-size:12px; background:#0A1F44;
    color:#dfe8ff; border-radius:10px; padding:14px 16px; line-height:1.55; overflow-x:auto;
  }}
  .cycle .mark {{ color:var(--cyan); font-weight:700; }}
  .pill-row {{ display:flex; flex-wrap:wrap; gap:8px; margin-top:10px; }}
  .pill {{ font-size:11.5px; font-weight:700; border-radius:999px; padding:5px 10px;
           background:#EAF1FF; color:var(--blue); border:1px solid #c9d8ff; }}
  .pill.em {{ background:#E3F8EF; color:#0C7C57; border-color:#c8ecd8; }}
  .pill.am {{ background:#FDF1DC; color:#9C6410; border-color:#f3dfb3; }}
  @media (max-width:800px) {{
    .kpis,.hl-grid {{ grid-template-columns:1fr; }}
  }}
</style>
</head>
<body>
  <div class="banner">
    <div class="brand"><i></i> INTELIX · IA-DLC · Plantilla base</div>
    <h1>Reporte de gestion — Iteraciones de Inception</h1>
    <p class="meta">Proyecto <strong>{esc(project)}</strong> · Generado {esc(generated)} ·
      Ledger <code>.aidlc/audit/inception-iterations.jsonl</code></p>
  </div>
<main>
  <section>
    <h2>Contexto del proyecto (ADN plantilla)</h2>
    <table>
      <thead><tr><th>Elemento</th><th>Valor</th></tr></thead>
      <tbody>{ctx_rows}</tbody>
    </table>
    <div class="pill-row">
      <span class="pill">Product-Definition</span>
      <span class="pill">Gate 0</span>
      <span class="pill em">develop = avance oficial</span>
      <span class="pill am">REWORK real vs 5% plan</span>
      <span class="pill">R1–R14 / estandares</span>
      <span class="pill">GitLab MR/issues</span>
    </div>
  </section>

  <h2 style="margin:4px 0 10px;font-size:16px">Elementos destacados de la plantilla</h2>
  <div class="hl-grid">{highlights}</div>

  <div class="kpis">
    <div class="kpi"><div class="n">{stats['inception_iterations']}</div><div class="l">Iteraciones Inception</div></div>
    <div class="kpi"><div class="n">{stats['construction_impacts']}</div><div class="l">Impactos Construction</div></div>
    <div class="kpi"><div class="n">{esc(fmt_hh(stats['hh_total']))}</div><div class="l">HH rework real (total)</div></div>
    <div class="kpi"><div class="n">{esc(fmt_hh(stats['hh_inception']))}</div><div class="l">HH en Inception</div></div>
    <div class="kpi"><div class="n">{esc(fmt_hh(stats['hh_construction']))}</div><div class="l">HH en Construction</div></div>
    <div class="kpi"><div class="n">{esc(stats['first_at'] or '—')} → {esc(stats['last_at'] or '—')}</div><div class="l">Ventana de fechas</div></div>
  </div>

  <section>
    <h2>Posicion en el ciclo AI-DLC</h2>
    <div class="cycle">Initialization → Ideation → <span class="mark">[Gate 0]</span> →
<span class="mark">Inception*</span> → <span class="mark">Construction*</span> → Operation
                Product-Definition          INC-ITER / SCOPE-RW (este reporte)</div>
  </section>

  <section>
    <h2>Taxonomia de disparadores</h2>
    <table>
      <thead><tr><th>Disparador</th><th>Eventos</th><th>HH</th></tr></thead>
      <tbody>{rows_kinds()}</tbody>
    </table>
  </section>

  <section>
    <h2>Linea de tiempo — Inception</h2>
    <table>
      <thead><tr><th>ID</th><th>Apertura</th><th>Cierre</th><th>Disparador</th><th>HH</th><th>Resumen</th><th>Evidencia</th></tr></thead>
      <tbody>{rows_inception()}</tbody>
    </table>
  </section>
  <section>
    <h2>Impacto en Construction y develop</h2>
    <p class="muted" style="margin:0 0 12px">Principio plantilla: el % oficial solo cuenta lo integrado en
      <code>develop</code> (<code>intelix-consolidado-avance</code>).</p>
    <table>
      <thead><tr><th>ID</th><th>Fecha</th><th>Devs</th><th>Ramas</th><th>HH</th><th>Efecto develop</th><th>Reabiertas</th><th>Nuevas</th></tr></thead>
      <tbody>{rows_impact()}</tbody>
    </table>
  </section>
  <section>
    <h2>Por desarrollador</h2>
    <ul>{dev_lis}</ul>
  </section>
  <section>
    <h2>Por rama feature</h2>
    <ul>{br_lis}</ul>
  </section>
  <section>
    <h2>Efectos acumulados en develop</h2>
    <ul>{dev_notes}</ul>
  </section>
  <section>
    <h2>Capacidades relacionadas</h2>
    <table>
      <thead><tr><th>Capacidad</th><th>Cruce</th></tr></thead>
      <tbody>
        <tr><td><code>intelix-estimacion</code></td><td>Baseline HH y REWORK_FACTOR planificado</td></tr>
        <tr><td><code>intelix-task-tracker</code></td><td>IDs reabiertas / nuevas / omitidas</td></tr>
        <tr><td><code>intelix-consolidado-avance</code></td><td>% oficial develop vs multi-rama</td></tr>
        <tr><td><code>intelix-certificacion</code></td><td>Dictamen de estandares pre-QA</td></tr>
        <tr><td>Gate 0 / Product-Definition</td><td>Contrato previo a Inception</td></tr>
      </tbody>
    </table>
  </section>
</main>
</body>
</html>
"""


def project_name() -> str:
    return load_project_context()["name"]


def write_reports(
    ledger: Path,
    out_md: Path,
    out_html: Path,
    name: str | None = None,
    ctx_overrides: dict[str, str] | None = None,
) -> dict[str, Any]:
    rows = load_ledger(ledger)
    stats = summarize(rows)
    generated = utc_now()
    ctx = load_project_context(ctx_overrides)
    if name:
        ctx["name"] = name
    pname = ctx["name"]
    out_md.parent.mkdir(parents=True, exist_ok=True)
    out_md.write_text(render_md(pname, stats, generated, ctx), encoding="utf-8")
    out_html.write_text(render_html(pname, stats, generated, ctx), encoding="utf-8")
    return stats


def cmd_report(args: argparse.Namespace) -> int:
    ledger = Path(args.ledger) if args.ledger else LEDGER
    out_md = Path(args.out_md) if args.out_md else REPORT_MD
    out_html = Path(args.out_html) if args.out_html else REPORT_HTML
    if not ledger.is_file():
        raise SystemExit(f"No hay ledger en {ledger}. Use `add` o `simulate`.")
    stats = write_reports(ledger, out_md, out_html)
    print(
        f"Reporte OK — Inception={stats['inception_iterations']} "
        f"impactos={stats['construction_impacts']} HH={stats['hh_total']:.2f}\n"
        f"  MD:   {out_md}\n"
        f"  HTML: {out_html}"
    )
    return 0


def patch_state_field(text: str, key: str, value: str) -> str:
    pattern = rf"(- \*\*{re.escape(key)}\*\*: `)([^`]*)(`)"
    if re.search(pattern, text):
        return re.sub(pattern, rf"\g<1>{value}\g<3>", text, count=1)
    return text


def cmd_sync_state(args: argparse.Namespace) -> int:
    ledger = Path(args.ledger) if getattr(args, "ledger", None) else LEDGER
    rows = load_ledger(ledger)
    stats = summarize(rows)
    if not STATE.is_file():
        raise SystemExit(f"No existe {STATE}")
    text = STATE.read_text(encoding="utf-8")
    updates = {
        "inception_iterations": str(stats["inception_iterations"]),
        "construction_impacts": str(stats["construction_impacts"]),
        "hh_impact_total": f"{stats['hh_total']:.2f}",
        "first_iteration_at": stats["first_at"],
        "last_iteration_at": stats["last_at"],
        "ledger_updated_at": utc_now(),
        "updated_at": today_iso(),
    }
    if "**inception_iterations**" not in text:
        block = (
            "\n## Inception iterations (resumen)\n\n"
            f"- **inception_iterations**: `{stats['inception_iterations']}`\n"
            f"- **construction_impacts**: `{stats['construction_impacts']}`\n"
            f"- **hh_impact_total**: `{stats['hh_total']:.2f}`\n"
            f"- **inception_gate_status**: `pending`\n"
            f"- **first_iteration_at**: `{stats['first_at']}`\n"
            f"- **last_iteration_at**: `{stats['last_at']}`\n"
            f"- **ledger_updated_at**: `{utc_now()}`\n\n"
        )
        text = re.sub(r"\n## Entry Mode Rules", block + "## Entry Mode Rules", text, count=1)
    else:
        if "**hh_impact_total**" not in text:
            text = text.replace(
                f"- **construction_impacts**: `{stats['construction_impacts']}`",
                f"- **construction_impacts**: `{stats['construction_impacts']}`\n"
                f"- **hh_impact_total**: `{stats['hh_total']:.2f}`",
            )
        for key, value in updates.items():
            text = patch_state_field(text, key, value)
    STATE.write_text(text, encoding="utf-8")
    print(
        f"aidlc-state: inception={stats['inception_iterations']} "
        f"impactos={stats['construction_impacts']} HH={stats['hh_total']:.2f}"
    )
    return 0


def _hh(estimated: float, by_task: dict[str, float] | None = None) -> dict[str, Any]:
    return {
        "estimated": estimated,
        "source": "simulation",
        "by_task": by_task or {},
        "skipped_hh": 0.0,
    }


SIM_EVENTS: list[dict[str, Any]] = [
    {
        "id": "INC-ITER-01",
        "phase": "inception",
        "kind": "gate_rejected",
        "opened_at": "2026-02-10",
        "closed_at": "2026-02-12",
        "requested_by": "Arquitecto",
        "approved_by": "Arquitecto",
        "summary": "Gate Inception rechazado: faltan historias medibles y STRIDE en auth.",
        "artifacts_changed": [
            "aidlc-docs/inception/requirements",
            "aidlc-docs/inception/application-design",
        ],
        "construction_impact": {
            "affected": False,
            "developers": [],
            "branches": [],
            "develop_effect": "",
            "tasks_reopened": [],
            "tasks_added": [],
            "tasks_skipped": [],
        },
        "hh": _hh(6.0),
        "evidence": {"issue": "#12", "mr": ""},
        "recorded_at": "2026-02-12T18:00:00Z",
    },
    {
        "id": "INC-ITER-02",
        "phase": "inception",
        "kind": "po_scope_change",
        "opened_at": "2026-02-18",
        "closed_at": "2026-02-20",
        "requested_by": "Product Owner",
        "approved_by": "Arquitecto",
        "summary": "PO agrega export CSV al modulo de reportes antes de cerrar Inception.",
        "artifacts_changed": [
            "Product-Definition/vision-document.md",
            "aidlc-docs/inception/requirements",
        ],
        "construction_impact": {
            "affected": False,
            "developers": [],
            "branches": [],
            "develop_effect": "",
            "tasks_reopened": [],
            "tasks_added": [],
            "tasks_skipped": [],
        },
        "hh": _hh(4.5),
        "evidence": {"issue": "#18", "mr": "!22"},
        "recorded_at": "2026-02-20T16:30:00Z",
    },
    {
        "id": "INC-ITER-03",
        "phase": "inception",
        "kind": "tech_discovery",
        "opened_at": "2026-02-24",
        "closed_at": "2026-02-25",
        "requested_by": "Desarrollador 1",
        "approved_by": "Arquitecto",
        "summary": "Hallazgo: API legacy no soporta paginacion; se ajusta diseno de listados.",
        "artifacts_changed": ["aidlc-docs/inception/application-design"],
        "construction_impact": {
            "affected": False,
            "developers": [],
            "branches": [],
            "develop_effect": "",
            "tasks_reopened": [],
            "tasks_added": [],
            "tasks_skipped": [],
        },
        "hh": _hh(3.0),
        "evidence": {"issue": "#21", "mr": ""},
        "recorded_at": "2026-02-25T20:00:00Z",
    },
    {
        "id": "SCOPE-RW-01",
        "phase": "construction_impact",
        "kind": "po_scope_change",
        "opened_at": "2026-03-08",
        "closed_at": "2026-03-11",
        "requested_by": "Product Owner",
        "approved_by": "Lider de proyecto",
        "summary": "Tras Gate Inception aprobado: PO pide MFA obligatorio; impacta auth ya en curso.",
        "artifacts_changed": [
            "Product-Definition/vision-document.md",
            "aidlc-docs/inception/requirements",
            "aidlc-docs/inception/application-design",
        ],
        "construction_impact": {
            "affected": True,
            "developers": ["Desarrollador 1", "Desarrollador 2"],
            "branches": ["feature/auth-login", "feature/auth-session"],
            "develop_effect": (
                "MR !40 (auth) revertido parcialmente en develop; "
                "rebase de 2 features; 2 stories reabiertas en indice de develop"
            ),
            "tasks_reopened": ["U1-T2", "U1-T4"],
            "tasks_added": ["U1-T7"],
            "tasks_skipped": [],
        },
        "hh": _hh(12.0, {"U1-T2": 4.0, "U1-T4": 3.5, "U1-T7": 4.5}),
        "evidence": {"issue": "#33", "mr": "!41"},
        "recorded_at": "2026-03-11T15:00:00Z",
    },
    {
        "id": "SCOPE-RW-02",
        "phase": "construction_impact",
        "kind": "reclassification",
        "opened_at": "2026-03-15",
        "closed_at": "2026-03-16",
        "requested_by": "Lider de proyecto",
        "approved_by": "Lider de proyecto",
        "summary": "Alcance crece: Intermedio → Mayor; se agrega gate de Operaciones.",
        "artifacts_changed": [
            "Product-Definition/ENTRY_MODE.md",
            ".aidlc/aidlc-state.md",
        ],
        "construction_impact": {
            "affected": True,
            "developers": ["Desarrollador 3", "Desarrollador 2"],
            "branches": ["feature/reportes-csv", "feature/api-listados"],
            "develop_effect": (
                "Plan y estimacion regenerados; develop mantiene avance oficial "
                "pero se congelan merges no relacionados hasta alinear historias"
            ),
            "tasks_reopened": [],
            "tasks_added": ["U3-T1", "U3-T2"],
            "tasks_skipped": ["U2-T6"],
        },
        "hh": _hh(8.5, {"U3-T1": 5.0, "U3-T2": 4.0, "-U2-T6": -0.5}),
        "evidence": {"issue": "#40", "mr": "!48"},
        "recorded_at": "2026-03-16T12:00:00Z",
    },
]


def cmd_simulate(_: argparse.Namespace) -> int:
    EXAMPLES.mkdir(parents=True, exist_ok=True)
    sim_ledger = EXAMPLES / "inception-iterations.jsonl"
    sim_ledger.write_text(
        "\n".join(json.dumps(e, ensure_ascii=False) for e in SIM_EVENTS) + "\n",
        encoding="utf-8",
    )
    (EXAMPLES / "hh-catalog.json").write_text(
        json.dumps(
            {
                "U1-T2": 4.0,
                "U1-T4": 3.5,
                "U1-T7": 4.5,
                "U3-T1": 5.0,
                "U3-T2": 4.0,
                "U2-T6": 0.5,
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    out_md = EXAMPLES / "inception-report.md"
    out_html = EXAMPLES / "inception-report.html"
    stats = summarize(SIM_EVENTS)
    generated = utc_now()
    name = "Demo INTELIX (simulacion multi-dev)"
    ctx = load_project_context(
        {
            "name": name,
            "entry_mode": "with-discovery",
            "classification": "intermedio → mayor (SCOPE-RW-02)",
            "gate_status": "approved",
            "inception_gate_status": "approved",
            "aidlc_version": "2.8.2",
            "has_product_definition": "si",
            "has_gate0_script": "si",
            "has_verificacion": "si",
            "has_consolidado_skill": "si",
        }
    )
    out_md.write_text(render_md(name, stats, generated, ctx), encoding="utf-8")
    out_html.write_text(render_html(name, stats, generated, ctx), encoding="utf-8")
    (EXAMPLES / "README.md").write_text(
        "\n".join(
            [
                "# Simulacion — iteraciones Inception + impacto HH",
                "",
                "Reporte alineado a elementos destacados de la plantilla INTELIX · IA-DLC",
                "(Gate 0, clasificacion, develop oficial, REWORK real, multi-dev).",
                "",
                "- **3** iteraciones de Inception (~13.5 HH)",
                "- **2** cambios con impacto Construction (~20.5 HH, multi-dev)",
                "- **34 HH** rework real acumulado",
                "",
                "```bash",
                "python .intelix/scripts/inception_iterations.py simulate",
                "```",
                "",
                "Abrir: `inception-report.html`",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(
        f"Simulacion lista en {EXAMPLES.relative_to(ROOT)}\n"
        f"  Iteraciones Inception: {stats['inception_iterations']} "
        f"({stats['hh_inception']} HH)\n"
        f"  Impactos Construction: {stats['construction_impacts']} "
        f"({stats['hh_construction']} HH)\n"
        f"  HH total: {stats['hh_total']}\n"
        f"  Abrir: {out_html}"
    )
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest="cmd", required=True)

    add = sub.add_parser("add", help="Registrar una iteracion o cambio de alcance")
    add.add_argument("--phase", required=True, choices=PHASES)
    add.add_argument("--kind", required=True, choices=KINDS)
    add.add_argument("--summary", required=True)
    add.add_argument("--id", default="")
    add.add_argument("--opened-at", default="")
    add.add_argument("--closed-at", default="")
    add.add_argument("--requested-by", default="")
    add.add_argument("--approved-by", default="")
    add.add_argument("--artifacts", default="")
    add.add_argument("--developers", default="")
    add.add_argument("--branches", default="")
    add.add_argument("--develop-effect", default="")
    add.add_argument("--tasks-reopened", default="")
    add.add_argument("--tasks-added", default="")
    add.add_argument("--tasks-skipped", default="")
    add.add_argument("--hh-delta", default="", help="HH de impacto (manual). Si se omite, se deriva del plan/catalogo")
    add.add_argument("--hh-catalog", default="", help="JSON opcional ID→HH")
    add.add_argument("--plan", default="", help="Ruta a PLAN-ESTIMACION-*-FINAL.xlsx")
    add.add_argument("--affects-construction", action="store_true")
    add.add_argument("--issue", default="")
    add.add_argument("--mr", default="")
    add.add_argument("--sync-state", action="store_true")
    add.set_defaults(func=cmd_add)

    lst = sub.add_parser("list", help="Listar eventos del ledger")
    lst.add_argument("--ledger", default="")
    lst.add_argument("--phase", choices=PHASES, default="")
    lst.set_defaults(func=cmd_list)

    st = sub.add_parser("status", help="Resumen de contadores e HH")
    st.add_argument("--ledger", default="")
    st.set_defaults(func=cmd_status)

    rep = sub.add_parser("report", help="Generar reporte MD + HTML")
    rep.add_argument("--ledger", default="")
    rep.add_argument("--out-md", default="")
    rep.add_argument("--out-html", default="")
    rep.set_defaults(func=cmd_report)

    syn = sub.add_parser("sync-state", help="Actualizar contadores en aidlc-state.md")
    syn.add_argument("--ledger", default="")
    syn.set_defaults(func=cmd_sync_state)

    sim = sub.add_parser("simulate", help="Generar ejemplo multi-dev + HH + reporte")
    sim.set_defaults(func=cmd_simulate)

    return p


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
