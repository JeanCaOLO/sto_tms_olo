#!/usr/bin/env python3
"""Lector del archivo de la verdad (SEGUIMIENTO-*.xlsx) para el flujo de commits.

Extrae de la hoja `Plan`:

- Columna B  `Nombre de tarea`  -> nombre del entregable / tarea
- Columna K  `ID Jerarquico`    -> ID estable (`PAR-002`, `PAR-002-003`)
- Columna H  `Recursos`         -> colaborador(es) asignado(s)
- Columna C  `% completado`     -> contexto de avance

Produce un catalogo `ID -> Nombre` seleccionable al momento de hacer commit,
en texto numerado (para elegir) o en JSON (para que lo consuma el hook Node,
que asi no necesita ninguna dependencia para leer .xlsx).

Fuente de verdad: `.intelix/outputs/seguimiento/SEGUIMIENTO-*.xlsx`
(se toma el snapshot mas reciente salvo que se indique `--archivo`).

Uso:
    python scripts/entregables/leer_entregables.py
    python scripts/entregables/leer_entregables.py --nivel tareas
    python scripts/entregables/leer_entregables.py --colaborador "Agente AI"
    python scripts/entregables/leer_entregables.py --entregable PAR-002
    python scripts/entregables/leer_entregables.py --json
    python scripts/entregables/leer_entregables.py --json --salida catalogo.json

Requiere: openpyxl>=3.1.0
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependencia externa
    sys.exit(
        "ERROR: falta openpyxl. Instalar con:\n"
        "  pip install \"openpyxl>=3.1.0\""
    )

# --------------------------------------------------------------------------
# Constantes del formato (ver docs/_revision/plan-implementacion-id-jerarquico.md)
# --------------------------------------------------------------------------

DIR_SEGUIMIENTO = Path(".intelix/outputs/seguimiento")
PREFIJO_ARCHIVO = "SEGUIMIENTO-"
HOJA_PLAN = "Plan"

# La cabecera NO esta en la fila 1: las filas 1-6 son titulo, codigo/lider,
# sello del corte, hitos de fase y leyenda. Se detecta, no se hardcodea.
FILA_CABECERA_ESPERADA = 7
MAX_FILAS_BUSQUEDA_CABECERA = 20

# Cabeceras buscadas (comparacion normalizada: sin acentos, minusculas).
# "ID Jerarquico" aparece sin acento en el archivo real, con acento en la
# documentacion: por eso se normaliza antes de comparar.
CAB_NOMBRE = ("nombre de tarea",)
CAB_ID = ("id jerarquico",)
CAB_RECURSOS = ("recursos", "nombres de los recursos")
CAB_AVANCE = ("% completado", "porcentaje completado")

# ID Jerarquico: entregable de 2 niveles, tarea de 3.
RE_ID_ENTREGABLE = re.compile(r"^[A-Z]{2,5}-\d{3}$")
RE_ID_TAREA = re.compile(r"^[A-Z]{2,5}-\d{3}-\d{3}$")

# Valores de la columna de recursos que NO son un colaborador real.
# Incluye la etiqueta "Recursos" que trae la fila de entregable.
RECURSOS_NO_INFORMATIVOS = {"", "-", "--", "recursos", "n/a", "na", "tbd"}

# Sufijo de dedicacion en las filas de tarea: "Senior 1[100%]".
RE_DEDICACION = re.compile(r"\[([^\]]*)\]")


# --------------------------------------------------------------------------
# Utilidades
# --------------------------------------------------------------------------


def normalizar(texto: object) -> str:
    """Minusculas sin acentos ni espacios sobrantes, para comparar cabeceras."""
    if texto is None:
        return ""
    plano = unicodedata.normalize("NFKD", str(texto))
    plano = "".join(c for c in plano if not unicodedata.combining(c))
    return " ".join(plano.lower().split())


def a_porcentaje(valor: object) -> int | None:
    """Normaliza el % completado, que puede venir como '100%', 1.0 o 100."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        num = float(valor)
        return int(round(num * 100)) if num <= 1 else int(round(num))
    texto = str(valor).strip().replace("%", "").replace(",", ".")
    try:
        return int(round(float(texto)))
    except ValueError:
        return None


def partir_recursos(valor: object) -> tuple[list[str], dict[str, str]]:
    """Separa la columna de recursos en colaboradores + su dedicacion.

    Soporta los dos formatos que produce la skill:
      - fila de entregable: "Senior 1; Senior 2"  o la etiqueta "Recursos"
      - fila de tarea:      "Senior 1[100%]"      o "Agente AI[100%]"
    """
    if valor is None:
        return [], {}

    nombres: list[str] = []
    dedicacion: dict[str, str] = {}

    for parte in re.split(r"[;\n]+", str(valor)):
        crudo = parte.strip()
        if not crudo:
            continue
        marca = RE_DEDICACION.search(crudo)
        nombre = RE_DEDICACION.sub("", crudo).strip(" ,;")
        if normalizar(nombre) in RECURSOS_NO_INFORMATIVOS:
            continue
        if nombre not in nombres:
            nombres.append(nombre)
        if marca and marca.group(1).strip():
            dedicacion[nombre] = marca.group(1).strip()

    return nombres, dedicacion


def coincide(aguja: str, valores: list[str]) -> bool:
    """Coincidencia por subcadena, insensible a acentos y mayusculas."""
    objetivo = normalizar(aguja)
    return any(objetivo in normalizar(v) for v in valores)


# --------------------------------------------------------------------------
# Modelo
# --------------------------------------------------------------------------


@dataclass
class Tarea:
    id: str
    nombre: str
    colaboradores: list[str] = field(default_factory=list)
    dedicacion: dict[str, str] = field(default_factory=dict)
    avance: int | None = None
    fila: int = 0

    def a_dict(self) -> dict:
        return {
            "id": self.id,
            "nombre": self.nombre,
            "colaboradores": self.colaboradores,
            "dedicacion": self.dedicacion,
            "avance": self.avance,
            "fila": self.fila,
        }


@dataclass
class Entregable:
    id: str
    nombre: str
    colaboradores: list[str] = field(default_factory=list)
    avance: int | None = None
    fila: int = 0
    tareas: list[Tarea] = field(default_factory=list)

    @property
    def prefijo(self) -> str:
        return self.id.split("-")[0]

    def a_dict(self) -> dict:
        return {
            "id": self.id,
            "nombre": self.nombre,
            "prefijo": self.prefijo,
            "colaboradores": self.colaboradores,
            "avance": self.avance,
            "fila": self.fila,
            "tareas": [t.a_dict() for t in self.tareas],
        }


# --------------------------------------------------------------------------
# Localizacion del archivo de la verdad
# --------------------------------------------------------------------------


def resolver_archivo(raiz: Path, explicito: str | None) -> Path:
    """Devuelve el SEGUIMIENTO-*.xlsx a leer (el mas reciente por defecto)."""
    carpeta = (raiz / DIR_SEGUIMIENTO).resolve()

    if explicito:
        ruta = Path(explicito)
        if not ruta.is_absolute():
            ruta = (raiz / ruta).resolve()
        else:
            ruta = ruta.resolve()
        # Se acepta solo dentro de la carpeta de seguimiento (evita traversal).
        if carpeta not in ruta.parents:
            raise SystemExit(
                f"ERROR: el archivo debe estar en {DIR_SEGUIMIENTO}\n"
                f"  recibido: {ruta}"
            )
        if not ruta.is_file():
            raise SystemExit(f"ERROR: no existe el archivo: {ruta}")
        return ruta

    if not carpeta.is_dir():
        raise SystemExit(
            f"ERROR: no existe la carpeta del archivo de la verdad:\n  {carpeta}\n"
            "Generar un snapshot con la skill intelix-seguimiento-plan."
        )

    candidatos = [
        p
        for p in carpeta.glob(f"{PREFIJO_ARCHIVO}*.xlsx")
        if not p.name.startswith("~$")  # descarta temporales de Excel
    ]
    if not candidatos:
        raise SystemExit(
            f"ERROR: no hay archivos {PREFIJO_ARCHIVO}*.xlsx en:\n  {carpeta}"
        )

    # El nombre lleva el corte al final (SEGUIMIENTO-<COD>-YYYYMMDD-HHMM). Se
    # ordena por ese timestamp, no alfabeticamente: con dos codigos de proyecto
    # distintos el orden por nombre elegiria el codigo "mayor", no el mas nuevo.
    candidatos.sort(key=lambda p: (_corte_del_nombre(p), p.stat().st_mtime))
    return candidatos[-1]


def _corte_del_nombre(ruta: Path) -> str:
    """Extrae 'YYYYMMDDHHMM' del nombre; cadena vacia si no sigue el patron."""
    marca = re.search(r"(\d{8})-(\d{4})$", ruta.stem)
    return f"{marca.group(1)}{marca.group(2)}" if marca else ""


# --------------------------------------------------------------------------
# Lectura de la hoja Plan
# --------------------------------------------------------------------------


def localizar_cabecera(hoja) -> tuple[int, dict[str, int]]:
    """Encuentra la fila de cabecera y el indice de cada columna que interesa.

    Se busca la fila que contenga 'ID Jerarquico' y 'Nombre de tarea'. Si no
    aparece, se cae a las posiciones documentadas (B, K, H, C) sobre la fila 7.
    """
    tope = min(hoja.max_row, MAX_FILAS_BUSQUEDA_CABECERA)

    for fila in range(1, tope + 1):
        etiquetas = {
            normalizar(hoja.cell(row=fila, column=col).value): col
            for col in range(1, hoja.max_column + 1)
        }
        col_id = next((etiquetas[c] for c in CAB_ID if c in etiquetas), None)
        col_nombre = next((etiquetas[c] for c in CAB_NOMBRE if c in etiquetas), None)
        if col_id and col_nombre:
            return fila, {
                "id": col_id,
                "nombre": col_nombre,
                "recursos": next(
                    (etiquetas[c] for c in CAB_RECURSOS if c in etiquetas), 8
                ),
                "avance": next(
                    (etiquetas[c] for c in CAB_AVANCE if c in etiquetas), 3
                ),
            }

    print(
        "AVISO: no se detecto la cabecera; se usan las posiciones por defecto "
        f"(fila {FILA_CABECERA_ESPERADA}: B=nombre, K=ID, H=recursos, C=avance).",
        file=sys.stderr,
    )
    return FILA_CABECERA_ESPERADA, {"nombre": 2, "id": 11, "recursos": 8, "avance": 3}


def leer_catalogo(ruta: Path) -> dict:
    """Construye el catalogo de entregables y tareas desde el .xlsx."""
    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    try:
        if HOJA_PLAN not in libro.sheetnames:
            raise SystemExit(
                f"ERROR: el archivo no tiene la hoja '{HOJA_PLAN}'.\n"
                f"  hojas: {libro.sheetnames}"
            )
        hoja = libro[HOJA_PLAN]
        fila_cab, cols = localizar_cabecera(hoja)

        entregables: list[Entregable] = []
        actual: Entregable | None = None
        huerfanas: list[str] = []
        ignoradas = 0

        for fila in range(fila_cab + 1, hoja.max_row + 1):
            crudo_id = hoja.cell(row=fila, column=cols["id"]).value
            crudo_nombre = hoja.cell(row=fila, column=cols["nombre"]).value

            # La sangria de la col B marca jerarquia en el Excel, pero aqui la
            # jerarquia se decide por el patron del ID (mas robusto).
            nombre = str(crudo_nombre).strip() if crudo_nombre is not None else ""
            ident = str(crudo_id).strip().upper() if crudo_id is not None else ""

            if not ident:
                # Filas sin ID: separadores y el 'TOTAL PROYECTO' del final.
                if nombre:
                    ignoradas += 1
                continue

            recursos, dedicacion = partir_recursos(
                hoja.cell(row=fila, column=cols["recursos"]).value
            )
            avance = a_porcentaje(hoja.cell(row=fila, column=cols["avance"]).value)

            if RE_ID_ENTREGABLE.match(ident):
                actual = Entregable(
                    id=ident,
                    nombre=nombre,
                    colaboradores=recursos,
                    avance=avance,
                    fila=fila,
                )
                entregables.append(actual)
            elif RE_ID_TAREA.match(ident):
                tarea = Tarea(
                    id=ident,
                    nombre=nombre,
                    colaboradores=recursos,
                    dedicacion=dedicacion,
                    avance=avance,
                    fila=fila,
                )
                padre = ident.rsplit("-", 1)[0]
                destino = actual if actual and actual.id == padre else None
                if destino is None:
                    destino = next((e for e in entregables if e.id == padre), None)
                if destino is None:
                    huerfanas.append(ident)
                    continue
                destino.tareas.append(tarea)
            else:
                ignoradas += 1

        # La fila de entregable puede traer solo la etiqueta "Recursos" en vez
        # de nombres: en ese caso el colaborador se deriva de sus tareas.
        for ent in entregables:
            if not ent.colaboradores:
                derivados: list[str] = []
                for tarea in ent.tareas:
                    for nombre_col in tarea.colaboradores:
                        if nombre_col not in derivados:
                            derivados.append(nombre_col)
                ent.colaboradores = derivados

        colaboradores = sorted(
            {c for e in entregables for c in e.colaboradores}
            | {c for e in entregables for t in e.tareas for c in t.colaboradores}
        )

        return {
            "generado": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
            "archivo_fuente": ruta.name,
            "ruta_fuente": str(ruta),
            "fila_cabecera": fila_cab,
            "columnas": {
                clave: openpyxl.utils.get_column_letter(idx)
                for clave, idx in cols.items()
            },
            "colaboradores": colaboradores,
            "entregables": [e.a_dict() for e in entregables],
            "totales": {
                "entregables": len(entregables),
                "tareas": sum(len(e.tareas) for e in entregables),
                "filas_ignoradas": ignoradas,
                "tareas_huerfanas": huerfanas,
            },
        }
    finally:
        libro.close()


# --------------------------------------------------------------------------
# Filtros y presentacion
# --------------------------------------------------------------------------


def filtrar(catalogo: dict, colaborador: str | None, entregable: str | None) -> dict:
    """Aplica los filtros de colaborador y de entregable sobre el catalogo."""
    salida = dict(catalogo)
    items = catalogo["entregables"]

    if entregable:
        objetivo = entregable.strip().upper()
        items = [e for e in items if e["id"] == objetivo]
        if not items:
            raise SystemExit(f"ERROR: no existe el entregable '{entregable}'.")

    if colaborador:
        recortados = []
        for ent in items:
            tareas = [
                t for t in ent["tareas"] if coincide(colaborador, t["colaboradores"])
            ]
            if tareas or coincide(colaborador, ent["colaboradores"]):
                copia = dict(ent)
                copia["tareas"] = tareas
                recortados.append(copia)
        items = recortados
        if not items:
            raise SystemExit(
                f"ERROR: ningun entregable o tarea asignado a '{colaborador}'.\n"
                f"  colaboradores en el archivo: {catalogo['colaboradores']}"
            )

    salida["entregables"] = items
    return salida


def _avance(valor: int | None) -> str:
    return "   —" if valor is None else f"{valor:3d}%"


def render_texto(catalogo: dict, nivel: str) -> str:
    """Lista numerada para elegir el ID al momento de commitear."""
    lineas: list[str] = []
    total_ent = catalogo["totales"]["entregables"]
    total_tar = catalogo["totales"]["tareas"]

    lineas.append(f"Archivo de la verdad : {catalogo['archivo_fuente']}")
    lineas.append(
        f"Cabecera             : fila {catalogo['fila_cabecera']} "
        f"(nombre={catalogo['columnas']['nombre']}, "
        f"id={catalogo['columnas']['id']}, "
        f"recursos={catalogo['columnas']['recursos']})"
    )
    lineas.append(f"Entregables/Tareas   : {total_ent} / {total_tar}")
    if catalogo["colaboradores"]:
        lineas.append(
            f"Colaboradores        : {', '.join(catalogo['colaboradores'])}"
        )
    lineas.append("")
    lineas.append(f"{'#':>4}  {'ID':<15} {'AVANCE':>6}  NOMBRE")
    lineas.append(f"{'':->4}  {'':-<15} {'':->6}  {'':-<52}")

    contador = 0
    for ent in catalogo["entregables"]:
        if nivel in ("entregables", "ambos"):
            contador += 1
            quien = ", ".join(ent["colaboradores"])
            lineas.append(
                f"{contador:>4}) {ent['id']:<15} {_avance(ent['avance'])}  "
                f"{ent['nombre']}"
            )
            if quien and nivel == "entregables":
                lineas.append(f"{'':>4}  {'':<15} {'':>6}  -> {quien}")
        if nivel in ("tareas", "ambos"):
            for tarea in ent["tareas"]:
                contador += 1
                quien = ", ".join(tarea["colaboradores"])
                sufijo = f"  [{quien}]" if quien else ""
                lineas.append(
                    f"{contador:>4}) {tarea['id']:<15} {_avance(tarea['avance'])}  "
                    f"  {tarea['nombre']}{sufijo}"
                )

    if contador == 0:
        lineas.append("  (sin resultados para los filtros indicados)")

    huerfanas = catalogo["totales"]["tareas_huerfanas"]
    if huerfanas:
        lineas.append("")
        lineas.append(f"AVISO: tareas sin entregable padre: {', '.join(huerfanas)}")

    return "\n".join(lineas)


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------


def construir_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Lee el archivo de la verdad SEGUIMIENTO-*.xlsx y lista los "
            "entregables/tareas (ID Jerarquico + Nombre) para elegirlos al commitear."
        )
    )
    parser.add_argument(
        "--archivo",
        help=f"SEGUIMIENTO-*.xlsx a leer (por defecto, el mas reciente de {DIR_SEGUIMIENTO}).",
    )
    parser.add_argument(
        "--raiz",
        default=".",
        help="Raiz del proyecto (por defecto, el directorio actual).",
    )
    parser.add_argument(
        "--nivel",
        choices=("entregables", "tareas", "ambos"),
        default="ambos",
        help="Que listar (por defecto: ambos).",
    )
    parser.add_argument(
        "--colaborador",
        help="Filtra por colaborador asignado (ej. \"Agente AI\", \"Senior 1\").",
    )
    parser.add_argument(
        "--entregable",
        help="Filtra por un entregable concreto (ej. PAR-002).",
    )
    parser.add_argument(
        "--json",
        action="store_true",
        dest="como_json",
        help="Emite el catalogo en JSON (para consumo por el hook de commit).",
    )
    parser.add_argument(
        "--salida",
        help="Escribe el JSON en un archivo en vez de stdout (implica --json).",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = construir_parser().parse_args(argv)
    raiz = Path(args.raiz).resolve()

    ruta = resolver_archivo(raiz, args.archivo)
    catalogo = leer_catalogo(ruta)
    catalogo = filtrar(catalogo, args.colaborador, args.entregable)

    if args.salida:
        destino = Path(args.salida)
        if not destino.is_absolute():
            destino = raiz / destino
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_text(
            json.dumps(catalogo, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(f"Catalogo escrito en: {destino}")
    elif args.como_json:
        print(json.dumps(catalogo, ensure_ascii=False, indent=2))
    else:
        print(render_texto(catalogo, args.nivel))

    return 0


if __name__ == "__main__":
    sys.exit(main())
