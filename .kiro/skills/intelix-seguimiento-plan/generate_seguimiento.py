# -*- coding: utf-8 -*-
"""
Generador de reporte de SEGUIMIENTO del plan definitivo.

Cruza dos fuentes de verdad:
  1. El plan de estimacion DEFINITIVO: PLAN-ESTIMACION-{CODIGO}-v{N}-FINAL.xlsx
     -> aporta el baseline: entregables, HH estimadas y fechas.
  2. El task-tracker: indice del registro AI-DLC activo, compatible con v1 y v2.
     -> aporta el avance real por tarea (estados con pesos) y por unidad.

Produce un snapshot inmutable versionado por fecha/hora:
  .intelix/outputs/seguimiento/SEGUIMIENTO-{CODIGO}-YYYYMMDD-HHMM.xlsx

La skill NO recalcula avance por su cuenta: consume los estados del tracker.
El -FINAL queda CONGELADO; este generador nunca lo modifica.

Uso:
  python .kiro/skills/intelix-seguimiento-plan/generate_seguimiento.py
  python .kiro/skills/intelix-seguimiento-plan/generate_seguimiento.py --code ABC12345
"""
import argparse
import json
import math
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ============================================================ CONFIG
# Raiz del proyecto: este script vive en .kiro/skills/intelix-seguimiento-plan/
ROOT = Path(__file__).resolve().parents[3]
DOCS_DIR = ROOT / "docs"

_COMMON_SCRIPTS = ROOT / ".intelix" / "scripts"
if not (_COMMON_SCRIPTS / "aidlc_paths.py").is_file():
    _COMMON_SCRIPTS = ROOT / "templates" / ".intelix" / "scripts"
sys.path.insert(0, str(_COMMON_SCRIPTS))
from aidlc_paths import resolve_layout  # noqa: E402

AI_DLC_LAYOUT = resolve_layout(ROOT)

# Nombre y codigo del proyecto NO se hardcodean: se leen de la configuracion
# central compartida por las skills (estimacion, seguimiento, task-tracker).
# Fuente unica de verdad: docs/scripts/project-config.json.
PROJECT_CONFIG_PATH = DOCS_DIR / "scripts" / "project-config.json"


def cargar_project_config(path=PROJECT_CONFIG_PATH):
    """Lee project-config.json. Devuelve dict; {} si no existe o esta corrupto."""
    p = Path(path)
    if not p.exists():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}


_PCFG = cargar_project_config()
# Valores por defecto SOLO como fallback inerte si el JSON no existe. El proyecto
# real los sustituye desde docs/scripts/project-config.json.
PROJECT_CODE = _PCFG.get("code", "PROYECTO")
PROJECT_NAME = _PCFG.get("name", "Proyecto")

# Columnas de la hoja Plan (heredadas del generador de estimacion):
#   9 = Inicio real, 10 = Fin real (las escribe este seguimiento),
#   11 = ID Jerarquico (nace en la estimacion; este seguimiento lo hereda).
COL_INICIO_REAL = 9   # I
COL_FIN_REAL = 10     # J
COL_ID_JERARQUICO = 11  # K
HEADER_ROW_PLAN = 4   # cabecera de la tabla en la hoja Plan del -FINAL
TRACKING_DIR = AI_DLC_LAYOUT.tracking_dir
PROGRESO_PATH = AI_DLC_LAYOUT.progress_index
AJUSTES_PATH = AI_DLC_LAYOUT.adjustments_index
AUDIT_PATHS = AI_DLC_LAYOUT.audit_files
OUTPUT_DIR = ROOT / ".intelix" / "outputs" / "seguimiento"

# Entregables de INCEPTION (su ventana E/F se toma del audit.md, fase Inception)
ENTREGABLES_INCEPTION = {
    "Mob Elaboration",
    "Modelo de la solucion",
    "Arquitectura de la solucion",
}

# Capacidad de equipo para re-anclar la construccion: 2 devs x 8 HH/dia = 16 HH/dia
HORAS_POR_DIA = 8.0
DEVS = 2
CAPACIDAD_DIA = HORAS_POR_DIA * DEVS


# ------------------------------------------------------------ CALENDARIO VE
# Replica minima del calendario habil de Venezuela usado por la skill de
# estimacion, para re-anclar la construccion desde el inicio real sin acoplar
# ambas skills por import.
def _pascua(anio):
    a = anio % 19
    b = anio // 100
    c = anio % 100
    d = b // 4
    e = b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4
    k = c % 4
    el = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * el) // 451
    mes = (h + el - 7 * m + 114) // 31
    dia = ((h + el - 7 * m + 114) % 31) + 1
    return date(anio, mes, dia)


def _feriados_ve(anio):
    p = _pascua(anio)
    return [
        date(anio, 1, 1), p - timedelta(days=48), p - timedelta(days=47),
        p - timedelta(days=3), p - timedelta(days=2), date(anio, 4, 19),
        date(anio, 5, 1), date(anio, 6, 24), date(anio, 7, 5),
        date(anio, 7, 24), date(anio, 10, 12), date(anio, 12, 24),
        date(anio, 12, 25), date(anio, 12, 31),
    ]


_FERIADOS = set()
for _a in (2026, 2027, 2028):
    _FERIADOS.update(_feriados_ve(_a))


def _es_laborable(d):
    return d.weekday() < 5 and d not in _FERIADOS


def _construir_dias_habiles(inicio, n=600):
    dias = []
    d = inicio
    while not _es_laborable(d):
        d += timedelta(days=1)
    while len(dias) < n:
        if _es_laborable(d):
            dias.append(d)
        d += timedelta(days=1)
    return dias

# ------------------------------------------------------------ pesos del tracker
# Identicos a los definidos por la skill task-tracker.
PESO_AGENTE = {
    "\u2b1c": 0.0,      # ⬜ Pendiente
    "\U0001f504": 0.0,  # 🔄 En progreso
    "\U0001f7e1": 0.9,  # 🟡 Completada (agente)
    "\U0001f527": 0.9,  # 🔧 En ajuste
    "\u2705": 1.0,      # ✅ Verificada
}
OMITIDA = "\u23ed"  # ⏭️ Omitida: fuera del denominador

ESTADO_NOMBRE = {
    "\u2b1c": "Pendiente",
    "\U0001f504": "En progreso",
    "\U0001f7e1": "Completada (agente)",
    "\U0001f527": "En ajuste",
    "\u2705": "Verificada",
    "\u23ed": "Omitida",
}

# Nombre EXACTO del entregable en el plan -> unidad del tracker. Vacio en la
# plantilla: cada proyecto mapea aqui sus entregables. Si un entregable no esta,
# el seguimiento no lo cruza con una unidad del tracker (degradacion elegante).
# Formato: { "<Nombre exacto del entregable en el plan>": "<id-unidad-tracker>" }
ENTREGABLE_A_UNIDAD = {}

# Entregables de inception ya completados (no rastreados por el tracker).
ENTREGABLE_OVERRIDE_PCT = {
    "Mob Elaboration": 100,
    "Modelo de la solucion": 100,
    "Arquitectura de la solucion": 100,
}

# ============================================================ ESTILOS (Intelix)
C_HEADER = "2E9CB6"
C_DARK = "1F4E79"
C_LIGHT = "BDD7EE"
C_GREY = "F2F2F2"
C_BORDER = "BFBFBF"

FONT_HEADER = Font(name="Segoe UI", bold=True, size=10, color="FFFFFF")
FONT_SUBHEADER = Font(name="Segoe UI", bold=True, size=10, color=C_DARK)
FONT_NORMAL = Font(name="Segoe UI", size=10, color="333333")
FONT_TITLE = Font(name="Segoe UI", bold=True, size=14, color=C_DARK)
FONT_TOTAL = Font(name="Segoe UI", bold=True, size=10, color="333333")

FILL_HEADER = PatternFill(start_color=C_HEADER, end_color=C_HEADER, fill_type="solid")
FILL_LIGHT = PatternFill(start_color=C_LIGHT, end_color=C_LIGHT, fill_type="solid")
FILL_GREY = PatternFill(start_color=C_GREY, end_color=C_GREY, fill_type="solid")

BORDER_THIN = Border(
    left=Side(style="thin", color=C_BORDER),
    right=Side(style="thin", color=C_BORDER),
    top=Side(style="thin", color=C_BORDER),
    bottom=Side(style="thin", color=C_BORDER),
)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


# ============================================================ PRECONDICION: plan -FINAL
FINAL_RE = re.compile(
    r"PLAN-ESTIMACION-(?P<code>.+?)-v(?P<n>\d+)-FINAL\.xlsx$", re.IGNORECASE
)


def find_plan_final(code=PROJECT_CODE, docs_dir=DOCS_DIR):
    """Devuelve (path, version) del plan definitivo de mayor version.

    Es una precondicion dura: sin plan -FINAL no tiene sentido dar seguimiento.
    """
    docs_dir = Path(docs_dir)
    candidatos = []
    if docs_dir.exists():
        for f in docs_dir.glob(f"PLAN-ESTIMACION-{code}-v*-FINAL.xlsx"):
            if f.name.startswith("~$"):
                continue
            m = FINAL_RE.search(f.name)
            if m:
                candidatos.append((int(m.group("n")), f))
    if not candidatos:
        return None, None
    candidatos.sort(key=lambda x: x[0])
    n, path = candidatos[-1]
    return path, f"v{n}"


# El generador de estimacion v3 anexa al nombre del entregable un sufijo de
# latencia de verificacion, ej. "Motor de Criterios  (+1.5d verif. humana)".
# Para mapear a la unidad del tracker (ENTREGABLE_A_UNIDAD) hay que quitar ese
# sufijo. Regex tolerante a 1+ espacios antes del parentesis.
_SUFIJO_VERIF_RE = re.compile(r"\s*\(\+[\d.,]+\s*d\s*verif\.?\s*humana\)\s*$", re.IGNORECASE)


def _nombre_base_entregable(nombre):
    """Nombre del entregable sin el sufijo '(+N.Nd verif. humana)'."""
    return _SUFIJO_VERIF_RE.sub("", str(nombre)).strip()


def leer_entregables_final(plan_path):
    """Lee la hoja Plan del -FINAL: entregables con HH estimadas y fechas.

    Devuelve lista ordenada de dicts:
      {nombre, hh, inicio, fin, pct_plan, id_jerarquico, task_ids}
    - `id_jerarquico`: ID del entregable en la col 11 (None si el plan es viejo).
    - `task_ids`: lista de IDs de tarea del plan (col 11), en orden. Vacia si el
      plan no trae la columna.
    Las filas de entregable no llevan sangria; las de tarea empiezan con espacios.
    """
    wb = openpyxl.load_workbook(plan_path, data_only=True)
    ws = wb["Plan"]
    entregables = []
    actual = None
    for r in range(5, ws.max_row + 1):
        nombre = ws.cell(r, 2).value
        if not nombre:
            continue
        texto = str(nombre)
        if texto.strip().upper().startswith("TOTAL"):
            actual = None
            continue
        id_col = ws.cell(r, COL_ID_JERARQUICO).value
        id_col = str(id_col).strip() if id_col not in (None, "") else None
        if texto.startswith(" "):
            # Fila de tarea: acumula su ID en el entregable actual
            if actual is not None:
                actual["task_ids"].append(id_col)
            continue
        # Fila de entregable
        actual = {
            "nombre": texto.strip(),
            "nombre_base": _nombre_base_entregable(texto.strip()),
            "hh": ws.cell(r, 4).value or 0,
            "inicio": ws.cell(r, 5).value or "",
            "fin": ws.cell(r, 6).value or "",
            "pct_plan": ws.cell(r, 3).value or "",
            "id_jerarquico": id_col,
            "task_ids": [],
        }
        entregables.append(actual)
    return entregables


# ============================================================ PARSER DEL TRACKER
def _peso_y_nombre_estado(estado_celda):
    """Dado el texto de la celda Estado, devuelve (simbolo, peso, es_verificada)."""
    if OMITIDA in estado_celda:
        return OMITIDA, None, False
    for simbolo, peso in PESO_AGENTE.items():
        if simbolo in estado_celda:
            return simbolo, peso, simbolo == "\u2705"
    return None, None, False


def _fmt_fecha_iso(valor):
    """Convierte '2026-08-26T10:16:51Z' -> '26/08/2026'. Vacio/—/sin-timestamp -> ''."""
    if not valor:
        return ""
    v = valor.strip()
    if v in ("\u2014", "-", "sin-timestamp"):
        return ""
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", v)
    if not m:
        return ""
    return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"


def _fecha_key(f):
    """Convierte 'DD/MM/AAAA' en tupla ordenable ('AAAA','MM','DD').

    Valores no fecha (vacio, 'En curso', '—') ordenan al final.
    """
    if not f or "/" not in str(f):
        return ("9999", "99", "99")
    d, m, a = str(f).split("/")
    return (a, m, d)


def _es_terminada(tarea):
    """Terminada = 🟡 / 🔧 / ✅ (peso 0.9 o 1.0)."""
    return tarea["peso"] is not None and tarea["peso"] >= 0.9


def fechas_reales_unidad(tareas):
    """Deriva (inicio_real, fin_real) de una unidad segun la regla del tracker.

    - Inicio = min(Inicio) de las tareas con timestamp valido.
    - Fin    = max(Fin) SOLO si todas las tareas (no omitidas) estan terminadas
               y tienen Fin valido; si no, 'En curso'.
    - Sin timestamps -> ('', '').
    """
    contadas = [t for t in tareas if t["peso"] is not None]  # excluye omitidas
    inicios = [_fmt_fecha_iso(t["inicio"]) for t in contadas if _fmt_fecha_iso(t["inicio"])]
    if not inicios:
        return "", ""

    inicio_real = min(inicios, key=_fecha_key)

    todas_terminadas = contadas and all(_es_terminada(t) for t in contadas)
    fines = [_fmt_fecha_iso(t["fin"]) for t in contadas if _fmt_fecha_iso(t["fin"])]
    if todas_terminadas and len(fines) == len(contadas):
        fin_real = max(fines, key=_fecha_key)
    else:
        fin_real = "En curso"
    return inicio_real, fin_real


def parse_progreso(path=PROGRESO_PATH):
    """Parsea INDICE-PROGRESO.md.

    Devuelve dict {unidad: {"tareas": [...], "plan": str, "agente": %, "verificado": %}}.
    Cada tarea: {id, nombre, simbolo, estado, peso, verificada, duracion, historias}.
    """
    path = Path(path)
    if not path.exists():
        return {}

    unidades = {}
    unidad = None
    plan_ref = ""

    for linea in path.read_text(encoding="utf-8").splitlines():
        s = linea.strip()
        if s.startswith("## Unidad:"):
            unidad = s.split("## Unidad:", 1)[1].strip()
            unidades[unidad] = {"tareas": [], "plan": "", "agente": 0, "verificado": 0}
            plan_ref = ""
            continue
        if unidad is None:
            continue
        if s.startswith("Plan:"):
            unidades[unidad]["plan"] = s.split("Plan:", 1)[1].strip().strip("`")
            continue
        if not s.startswith("|"):
            continue
        celdas = [c.strip() for c in s.strip("|").split("|")]
        if len(celdas) < 3:
            continue
        if celdas[0].upper() == "ID" or set(celdas[0]) <= {"-", ":"}:
            continue  # cabecera o separador
        simbolo, peso, verificada = _peso_y_nombre_estado(celdas[2])
        tarea = {
            "id": celdas[0],
            "nombre": celdas[1],
            "simbolo": simbolo or "",
            "estado": ESTADO_NOMBRE.get(simbolo, celdas[2]),
            "peso": peso,
            "verificada": verificada,
            "inicio": celdas[3] if len(celdas) > 3 else "",
            "fin": celdas[4] if len(celdas) > 4 else "",
            "duracion": celdas[5] if len(celdas) > 5 else "",
            "historias": celdas[7] if len(celdas) > 7 else "",
        }
        unidades[unidad]["tareas"].append(tarea)

    # Calcular % agente y % verificado por unidad (excluye omitidas del denominador)
    for u, data in unidades.items():
        contadas = [t for t in data["tareas"] if t["peso"] is not None]
        n = len(contadas)
        if n == 0:
            data["agente"] = 0
            data["verificado"] = 0
        else:
            data["agente"] = round(sum(t["peso"] for t in contadas) / n * 100)
            data["verificado"] = round(sum(1 for t in contadas if t["verificada"]) / n * 100)
    return unidades


def parse_ajustes(path=AJUSTES_PATH):
    """Parsea INDICE-AJUSTES.md. Devuelve lista de dicts (vacia si no hay)."""
    path = Path(path)
    if not path.exists():
        return []
    ajustes = []
    for linea in path.read_text(encoding="utf-8").splitlines():
        s = linea.strip()
        if not s.startswith("|"):
            continue
        celdas = [c.strip() for c in s.strip("|").split("|")]
        if len(celdas) < 4:
            continue
        if celdas[0].upper() == "ID" or set(celdas[0]) <= {"-", ":"}:
            continue
        if celdas[0] in ("", "\u2014"):  # placeholder "sin ajustes"
            continue
        ajustes.append({
            "id": celdas[0],
            "origen": celdas[1],
            "descripcion": celdas[2],
            "estado": celdas[3],
            "inicio": celdas[4] if len(celdas) > 4 else "",
            "fin": celdas[5] if len(celdas) > 5 else "",
            "duracion": celdas[6] if len(celdas) > 6 else "",
        })
    return ajustes


def parse_hitos_fase(paths=AUDIT_PATHS):
    """Deriva del audit de v1 o sus shards de v2 la ventana de INCEPTION.

    Recorre los headers '## <titulo>' y sus '**Timestamp**: <ISO>'.
    - Inicio Inception = primer timestamp del audit.
    - Fin Inception    = ultimo timestamp ANTES del primer header que contenga
      'CONSTRUCTION' (case-insensitive). El inicio real de construccion de
      CODIGO no sale de aqui: viene del task-tracker.
    Devuelve {'inicio_inception': 'DD/MM/AAAA', 'fin_inception': 'DD/MM/AAAA'}
    o {} si no hay datos.
    """
    if isinstance(paths, (str, Path)):
        paths = (Path(paths),)
    paths = tuple(Path(path) for path in paths if Path(path).is_file())
    if not paths:
        return {}

    ts_re = re.compile(r"\*\*Timestamp\*\*:\s*([0-9T:\-Z]+)")
    eventos = []
    for path in paths:
        en_construccion = False
        for linea in path.read_text(encoding="utf-8").splitlines():
            s = linea.strip()
            if s.startswith("##") and "CONSTRUCTION" in s.upper():
                en_construccion = True
            m = ts_re.search(s)
            if m:
                eventos.append((m.group(1), en_construccion))

    if not eventos:
        return {}
    eventos.sort(key=lambda evento: evento[0])
    primero = eventos[0][0]
    timestamps_inception = [ts for ts, construction in eventos if not construction]
    ultimo_inception = timestamps_inception[-1] if timestamps_inception else primero
    return {
        "inicio_inception": _fmt_fecha_iso(primero),
        "fin_inception": _fmt_fecha_iso(ultimo_inception),
    }


def _fecha_a_ddmmaaaa(d):
    return d.strftime("%d/%m/%Y")


def _parse_ddmmaaaa(s):
    """'DD/MM/AAAA' -> date, o None si no es fecha valida."""
    if not s or "/" not in str(s):
        return None
    try:
        d, m, a = str(s).split("/")
        return date(int(a), int(m), int(d))
    except (ValueError, TypeError):
        return None


def reanclar_construccion(filas, fecha_inicio_real):
    """Re-programa las fechas E/F estimadas de la CONSTRUCCION desde el inicio real.

    En vez de mantener las fechas estimadas originales del plan (que arrancaban
    antes del inicio real de construccion), se re-anclan al `fecha_inicio_real`
    (date) y se encadenan secuencialmente por HH estimadas sobre dias habiles VE
    a 16 HH/dia. Se aplica a todos los entregables que NO son de inception,
    empezando por el primero de construccion.

    - Inception: no se toca (conserva su ventana real del audit).
    - Unidad ya iniciada (con inicio_real): usa su fecha real; las siguientes
      se encadenan a partir del fin re-anclado de esa unidad.
    - `comienzo_ef` / `fin_ef` se sobrescriben con las fechas re-ancladas
      (salvo que la unidad tenga fin real, que prevalece).

    Modifica las filas in situ.
    """
    dias = _construir_dias_habiles(fecha_inicio_real)

    def fecha_en_offset(h):
        idx = int(h // CAPACIDAD_DIA)
        return dias[min(idx, len(dias) - 1)]

    def fecha_fin_offset(h_fin):
        idx = int(max(h_fin - 1e-9, 0.0) // CAPACIDAD_DIA)
        return dias[min(idx, len(dias) - 1)]

    cursor = 0.0  # horas acumuladas de equipo desde el inicio real
    for f in filas:
        if f["nombre"] in ENTREGABLES_INCEPTION:
            continue  # inception conserva su ventana real
        hh = f.get("hh") or 0

        comienzo_reanclado = fecha_en_offset(cursor)
        fin_reanclado = fecha_fin_offset(cursor + hh) if hh else comienzo_reanclado

        # Si la unidad ya termino de verdad, respetar su fin real.
        fin_real = f.get("fin_real")
        if fin_real and fin_real not in ("", "En curso"):
            f["comienzo_ef"] = f.get("inicio_real") or _fecha_a_ddmmaaaa(comienzo_reanclado)
            f["fin_ef"] = fin_real
        else:
            # Unidad en curso o pendiente: comienzo real si existe, si no re-anclado
            f["comienzo_ef"] = f.get("inicio_real") or _fecha_a_ddmmaaaa(comienzo_reanclado)
            f["fin_ef"] = _fecha_a_ddmmaaaa(fin_reanclado)

        cursor += hh


def parse_us(path=PROGRESO_PATH):
    """Extrae la tabla 'Avance por Historia de Usuario' entre los marcadores."""
    path = Path(path)
    if not path.exists():
        return []
    contenido = path.read_text(encoding="utf-8")
    ini = contenido.find("US-SUMMARY-BEGIN")
    fin = contenido.find("US-SUMMARY-END")
    if ini == -1 or fin == -1:
        return []
    bloque = contenido[ini:fin]
    filas = []
    for linea in bloque.splitlines():
        s = linea.strip()
        if not s.startswith("|"):
            continue
        celdas = [c.strip() for c in s.strip("|").split("|")]
        if len(celdas) < 6:
            continue
        if celdas[0].lower() == "historia" or set(celdas[0]) <= {"-", ":"}:
            continue
        filas.append(celdas)  # [Historia, Tareas, Agente, Verificado, Inicio, Fin]
    return filas


# ============================================================ CRUCE
def cruzar(entregables_final, unidades, hitos=None):
    """Une cada entregable del -FINAL con su unidad del tracker.

    Calcula, ademas del avance, las fechas EFECTIVAS de Comienzo/Fin (E/F) que
    reemplazaran a las estimadas del plan:
      - Entregables de INCEPTION: ventana real de Inception (audit.md).
      - Entregables de CONSTRUCCION rastreados: Comienzo = inicio real del
        codigo (tracker); Fin = fin real (tracker) o, si no terminaron, la
        fecha ESTIMADA del plan para esa unidad (regla 4=b).
      - Entregables sin unidad ni inception: se conservan las fechas del plan.
    Conserva tambien inicio_real/fin_real por unidad para las columnas I/J.
    """
    hitos = hitos or {}
    filas = []
    for ent in entregables_final:
        nombre = ent["nombre"]
        # Nombre sin sufijo "(+N.Nd verif. humana)" para mapear a la unidad
        nombre_base = ent.get("nombre_base") or _nombre_base_entregable(nombre)
        unidad = ENTREGABLE_A_UNIDAD.get(nombre_base)
        est_inicio = ent["inicio"]  # fecha estimada del plan (DD/MM/AAAA)
        est_fin = ent["fin"]

        if unidad and unidad in unidades:
            agente = unidades[unidad]["agente"]
            verificado = unidades[unidad]["verificado"]
            tareas = unidades[unidad]["tareas"]
            inicio_real, fin_real = fechas_reales_unidad(tareas)
            # E = inicio real si existe, si no la estimada
            comienzo_ef = inicio_real or est_inicio
            # F = fin real si la unidad termino; si "En curso" o vacio -> estimada
            if fin_real and fin_real != "En curso":
                fin_ef = fin_real
            else:
                # Unidad no terminada: fin estimado del plan (regla 4=b), pero
                # nunca anterior al comienzo real (evita fin < inicio incoherente).
                if inicio_real and _fecha_key(est_fin) < _fecha_key(inicio_real):
                    fin_ef = "En curso"
                else:
                    fin_ef = est_fin
        elif nombre_base in ENTREGABLE_OVERRIDE_PCT:
            agente = ENTREGABLE_OVERRIDE_PCT[nombre_base]
            verificado = 0
            tareas = []
            inicio_real, fin_real = "", ""
            # Inception: ventana real del audit.md (con fallback a lo estimado)
            comienzo_ef = hitos.get("inicio_inception") or est_inicio
            fin_ef = hitos.get("fin_inception") or est_fin
        else:
            agente = 0
            verificado = 0
            tareas = []
            inicio_real, fin_real = "", ""
            comienzo_ef = est_inicio
            fin_ef = est_fin

        filas.append({
            **ent,
            "unidad": unidad or "",
            "agente": agente,
            "verificado": verificado,
            "tareas": tareas,
            "inicio_real": inicio_real,
            "fin_real": fin_real,
            "comienzo_ef": comienzo_ef,
            "fin_ef": fin_ef,
        })
    return filas


def kpis_globales(unidades):
    """% agente y % verificado global sobre todas las tareas no omitidas."""
    contadas = []
    for data in unidades.values():
        contadas.extend(t for t in data["tareas"] if t["peso"] is not None)
    n = len(contadas)
    if n == 0:
        return 0, 0, 0
    agente = round(sum(t["peso"] for t in contadas) / n * 100)
    verificado = round(sum(1 for t in contadas if t["verificada"]) / n * 100)
    return agente, verificado, n


# ============================================================ EXCEL
def _header_row(ws, row, headers):
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=ci, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN


def _style_row(ws, row, cols, font, fill=None, left_cols=(2,)):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = font
        if fill:
            cell.fill = fill
        cell.border = BORDER_THIN
        cell.alignment = ALIGN_LEFT if c in left_cols else ALIGN_CENTER


def write_resumen(ws, filas, ag_global, ver_global, n_tareas, meta):
    ws.merge_cells("A1:F1")
    ws["A1"] = f"Seguimiento del Plan — {PROJECT_CODE} ({meta['plan_version']} FINAL)"
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = ALIGN_LEFT

    info = [
        ("Proyecto", PROJECT_NAME),
        ("Codigo", PROJECT_CODE),
        ("Plan base (definitivo)", meta["plan_file"]),
        ("Fecha/hora del corte", meta["snapshot"]),
        ("Tareas rastreadas", str(n_tareas)),
        ("% Avance agente (global)", f"{ag_global}%"),
        ("% Verificado (global)", f"{ver_global}%"),
    ]
    row = 3
    for label, value in info:
        ws.cell(row=row, column=1, value=label).font = FONT_SUBHEADER
        ws.cell(row=row, column=2, value=value).font = FONT_NORMAL
        row += 1

    row += 1
    headers = ["Entregable", "HH estimadas", "% Agente", "% Verificado", "Unidad tracker", "Estado"]
    _header_row(ws, row, headers)
    row += 1

    total_hh = 0
    for f in filas:
        estado = "Sin rastrear" if not f["unidad"] and f["agente"] == 0 else (
            "Inception (100%)" if not f["unidad"] else "Rastreado")
        ws.cell(row=row, column=1, value=f["nombre"])
        ws.cell(row=row, column=2, value=f["hh"])
        ws.cell(row=row, column=3, value=f"{f['agente']}%")
        ws.cell(row=row, column=4, value=f"{f['verificado']}%")
        ws.cell(row=row, column=5, value=f["unidad"] or "—")
        ws.cell(row=row, column=6, value=estado)
        _style_row(ws, row, 6, FONT_NORMAL, left_cols=(1,))
        total_hh += f["hh"] or 0
        row += 1

    ws.cell(row=row, column=1, value="TOTAL")
    ws.cell(row=row, column=2, value=total_hh)
    ws.cell(row=row, column=3, value=f"{ag_global}%")
    ws.cell(row=row, column=4, value=f"{ver_global}%")
    _style_row(ws, row, 6, FONT_TOTAL, FILL_LIGHT, left_cols=(1,))

    for i, w in enumerate([42, 14, 12, 14, 22, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _pct_tarea(tarea):
    """% de una tarea del tracker segun su peso: 0 / 90 / 100."""
    if tarea["peso"] is None:  # omitida
        return None
    return int(round(tarea["peso"] * 100))


def actualizar_plan_en_copia(ws, entregables):
    """Reescribe la columna C (% completado) de la hoja Plan copiada del -FINAL
    y añade dos columnas: 'Inicio real' y 'Fin real' desde el task-tracker.

    Conserva intacto el formato/estructura original. Las columnas Comienzo/Fin
    reflejan fechas REALES tanto en entregables como en tareas; las columnas
    I/J (Inicio/Fin real) repiten ese detalle por tarea/unidad del tracker.

    - % completado: por entregable (agregado) y por tarea (posicion N<->N).
    - Comienzo/Fin (E/F):
        * Entregable: fechas efectivas de fase/unidad (ver cruzar()).
        * Tarea: inicio/fin real de esa tarea del tracker (posicion N<->N).
          Las tareas del plan sin correspondencia en el tracker muestran '—'.

    `entregables` es la lista de dicts producida por cruzar().
    """
    # Cabecera de las dos columnas nuevas, con el estilo de las demas cabeceras.
    for col, titulo in ((COL_INICIO_REAL, "Inicio real"), (COL_FIN_REAL, "Fin real")):
        cell = ws.cell(HEADER_ROW_PLAN, col, value=titulo)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN
    ws.column_dimensions[get_column_letter(COL_INICIO_REAL)].width = 13
    ws.column_dimensions[get_column_letter(COL_FIN_REAL)].width = 13

    por_nombre = {e["nombre"]: e for e in entregables}

    ent_actual = None
    idx_tarea = 0
    id_map = {}          # id_jerarquico_tarea -> tarea del tracker (join por ID)
    usar_id = False      # True si el entregable actual y el plan traen IDs
    for r in range(5, ws.max_row + 1):
        nombre = ws.cell(r, 2).value
        if nombre is None:
            continue
        texto = str(nombre)
        es_tarea = texto.startswith(" ")
        es_total = texto.strip().upper().startswith("TOTAL")
        if es_total:
            ent_actual = None
            continue

        if not es_tarea:
            # Fila de entregable
            ent_actual = por_nombre.get(texto.strip())
            idx_tarea = 0
            id_map = {}
            usar_id = False
            if ent_actual is not None:
                ws.cell(r, 3, value=f"{ent_actual['agente']}%")
                # E/F: fechas efectivas (reales de fase/tracker, fin estimado si no terminó)
                ws.cell(r, 5, value=ent_actual["comienzo_ef"])
                ws.cell(r, 6, value=ent_actual["fin_ef"])
                ws.cell(r, COL_INICIO_REAL, value=ent_actual["inicio_real"] or "—")
                ws.cell(r, COL_FIN_REAL, value=ent_actual["fin_real"] or "—")
                # Preparar el JOIN POR ID JERARQUICO: la tarea i (0-based) de la
                # unidad recibe id = "<ID entregable plan>-<i+1:03d>". Solo si el
                # plan trae ID de entregable (col 11); si no, se usa posicion.
                id_ent = ent_actual.get("id_jerarquico")
                if id_ent:
                    usar_id = True
                    for i, t in enumerate(ent_actual["tareas"]):
                        id_map[f"{id_ent}-{i + 1:03d}"] = t
            continue

        # Fila de tarea: cruzar por ID Jerarquico si el plan lo trae; si no, por posicion
        if ent_actual is None:
            continue
        tareas = ent_actual["tareas"]
        t = None
        if usar_id:
            id_tarea_plan = ws.cell(r, COL_ID_JERARQUICO).value
            id_tarea_plan = str(id_tarea_plan).strip() if id_tarea_plan not in (None, "") else None
            t = id_map.get(id_tarea_plan)  # None si es doc/gestion sin match
        elif idx_tarea < len(tareas):
            # Fallback: mapeo por posicion (planes viejos sin columna ID)
            t = tareas[idx_tarea]

        if t is not None:
            pct = _pct_tarea(t)
            ini = _fmt_fecha_iso(t["inicio"]) or "—"
            fin = _fmt_fecha_iso(t["fin"]) or "—"
            ws.cell(r, 3, value=f"{pct}%" if pct is not None else "N/A")
            # Comienzo/Fin (E/F): fechas REALES de la tarea
            ws.cell(r, 5, value=ini)
            ws.cell(r, 6, value=fin)
            # I/J: mismo detalle real
            ws.cell(r, COL_INICIO_REAL, value=ini)
            ws.cell(r, COL_FIN_REAL, value=fin)
        else:
            # El plan tiene mas tareas que el tracker (doc, gestion, etc.)
            # o el entregable no esta rastreado -> hereda 0% salvo override.
            if not ent_actual["unidad"] and ent_actual["agente"] == 100:
                ws.cell(r, 3, value="100%")
            else:
                ws.cell(r, 3, value="0%")
            # Sin correspondencia en el tracker: sin fecha real
            ws.cell(r, 5, value="—")
            ws.cell(r, 6, value="—")
            ws.cell(r, COL_INICIO_REAL, value="—")
            ws.cell(r, COL_FIN_REAL, value="—")
        idx_tarea += 1


def _sello_seguimiento(ws, meta, hitos=None):
    """Añade a la hoja Plan copiada un sello de seguimiento + leyenda del %.

    Inserta 3 filas antes de la cabecera de la tabla (para no alterar el bloque
    de datos ni sus estilos): corte + plan base, hitos de fase, y leyenda 0/90/100.
    """
    hitos = hitos or {}
    ws.insert_rows(3, amount=3)
    # 11 cols: 8 originales + Inicio real (9) + Fin real (10) + ID Jerarquico (11)
    ancho = 11

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ancho)
    c3 = ws.cell(3, 1, value=(
        f"SEGUIMIENTO — Corte: {meta['snapshot']}  |  "
        f"Plan base: {meta['plan_file']} ({meta['plan_version']} FINAL)"
    ))
    c3.font = FONT_SUBHEADER
    c3.alignment = ALIGN_LEFT

    ini_inc = hitos.get("inicio_inception", "—")
    fin_inc = hitos.get("fin_inception", "—")
    ini_cons = meta.get("inicio_construccion", "—")
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=ancho)
    c4 = ws.cell(4, 1, value=(
        f"Hitos de fase:  Inception {ini_inc} - {fin_inc}  |  "
        f"Construccion (codigo) desde {ini_cons}  |  Construccion: en curso"
    ))
    c4.font = FONT_SUBHEADER
    c4.alignment = ALIGN_LEFT

    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=ancho)
    c5 = ws.cell(5, 1, value=(
        "Leyenda % completado:  0% = tarea no realizada;  "
        "90% = tarea realizada por el agente, pendiente de verificacion humana;  "
        "100% = tarea realizada por el agente y verificada por una persona.  "
        "Columnas Comienzo/Fin = fechas reales de fase/hito; Inicio/Fin real = detalle por tarea."
    ))
    c5.font = FONT_NORMAL
    c5.alignment = ALIGN_LEFT


def write_ajustes(ws, ajustes):
    ws.merge_cells("A1:G1")
    ws["A1"] = "Ajustes Post-Entrega"
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = ALIGN_LEFT

    headers = ["ID", "Origen", "Descripcion", "Estado", "Inicio", "Fin", "Duracion"]
    row = 3
    _header_row(ws, row, headers)
    row += 1

    if not ajustes:
        ws.cell(row=row, column=1, value="Sin ajustes registrados a la fecha del corte.")
        ws.cell(row=row, column=1).font = FONT_NORMAL
    else:
        for a in ajustes:
            ws.cell(row=row, column=1, value=a["id"])
            ws.cell(row=row, column=2, value=a["origen"])
            ws.cell(row=row, column=3, value=a["descripcion"])
            ws.cell(row=row, column=4, value=a["estado"])
            ws.cell(row=row, column=5, value=a["inicio"])
            ws.cell(row=row, column=6, value=a["fin"])
            ws.cell(row=row, column=7, value=a["duracion"])
            _style_row(ws, row, 7, FONT_NORMAL, left_cols=(2, 3))
            row += 1

    for i, w in enumerate([10, 14, 50, 14, 20, 20, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_us(ws, us_filas):
    ws.merge_cells("A1:F1")
    ws["A1"] = "Avance por Historia de Usuario"
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = ALIGN_LEFT

    headers = ["Historia", "Tareas", "% Agente", "% Verificado", "Inicio", "Fin"]
    row = 3
    _header_row(ws, row, headers)
    row += 1

    if not us_filas:
        ws.cell(row=row, column=1, value="Tabla de historias no disponible en el indice.")
        ws.cell(row=row, column=1).font = FONT_NORMAL
    else:
        for celdas in us_filas:
            for ci in range(6):
                ws.cell(row=row, column=ci + 1, value=celdas[ci] if ci < len(celdas) else "")
            _style_row(ws, row, 6, FONT_NORMAL, left_cols=(1,))
            row += 1

    for i, w in enumerate([12, 10, 12, 14, 22, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ============================================================ MAIN
def parse_args(argv=None):
    p = argparse.ArgumentParser(
        description="Genera un snapshot de seguimiento del plan definitivo, "
                    "versionado por fecha/hora, cruzando el -FINAL con el task-tracker."
    )
    p.add_argument("--code", default=PROJECT_CODE,
                   help="Codigo de proyecto (por defecto, el de project-config.json)")
    return p.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)

    plan_path, plan_version = find_plan_final(args.code)
    if plan_path is None:
        print("=" * 60)
        print("  ERROR: no se encontro un plan DEFINITIVO.")
        print(f"  Se requiere docs/PLAN-ESTIMACION-{args.code}-vN-FINAL.xlsx")
        print("  Marca primero la estimacion como definitiva (--final).")
        print("=" * 60)
        return 2

    entregables_final = leer_entregables_final(plan_path)
    unidades = parse_progreso()
    ajustes = parse_ajustes()
    us_filas = parse_us()
    hitos = parse_hitos_fase()

    filas = cruzar(entregables_final, unidades, hitos)
    ag_global, ver_global, n_tareas = kpis_globales(unidades)

    # Inicio real de construccion de CODIGO = min inicio_real de las unidades
    inicios_codigo = [f["inicio_real"] for f in filas if f["unidad"] and f["inicio_real"]]
    inicio_construccion = min(inicios_codigo, key=_fecha_key) if inicios_codigo else "—"

    # Re-anclar la construccion desde el inicio real (no estimar antes de empezar)
    base_reancla = _parse_ddmmaaaa(inicio_construccion)
    if base_reancla is not None:
        reanclar_construccion(filas, base_reancla)

    ahora = datetime.now()
    snapshot = ahora.strftime("%d/%m/%Y %H:%M")
    stamp = ahora.strftime("%Y%m%d-%H%M")
    meta = {
        "plan_file": plan_path.name,
        "plan_version": plan_version,
        "snapshot": snapshot,
        "inicio_construccion": inicio_construccion,
    }

    # Partir del -FINAL para conservar su formato exacto (4 hojas y estilos).
    wb = openpyxl.load_workbook(plan_path)

    # Reescribir la columna C (% completado) de la hoja Plan con el avance real.
    ws_plan = wb["Plan"]
    actualizar_plan_en_copia(ws_plan, filas)
    _sello_seguimiento(ws_plan, meta, hitos)

    # Anexar hojas de seguimiento como valor extra.
    write_resumen(wb.create_sheet("Seguimiento"), filas, ag_global, ver_global, n_tareas, meta)
    write_ajustes(wb.create_sheet("Ajustes"), ajustes)
    write_us(wb.create_sheet("Avance US"), us_filas)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    filename = f"SEGUIMIENTO-{args.code}-{stamp}.xlsx"
    output_path = OUTPUT_DIR / filename
    wb.save(str(output_path))

    print("=" * 60)
    print(f"  {filename} generado")
    print("=" * 60)
    print(f"  Plan base:        {plan_path.name} ({plan_version} FINAL)")
    print(f"  Corte:            {snapshot}")
    print(f"  Tareas:           {n_tareas}")
    print(f"  % Agente global:  {ag_global}%")
    print(f"  % Verificado:     {ver_global}%")
    print("-" * 60)
    for f in filas:
        print(f"    - {f['nombre']:<44} ag {f['agente']}% / ver {f['verificado']}%")
    print("=" * 60)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
