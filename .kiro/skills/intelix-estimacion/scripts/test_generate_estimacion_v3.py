import importlib.util
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path


SCRIPT_PATH = Path(__file__).with_name("generate_estimacion_v3.py")
SPEC = importlib.util.spec_from_file_location("generate_estimacion_v3", SCRIPT_PATH)
estimacion = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(estimacion)


class EntradasTest(unittest.TestCase):
    def test_normaliza_dedicacion_fraccion_porcentaje_y_texto(self):
        self.assertEqual(estimacion.normalizar_dedicacion(0.8), 0.8)
        self.assertEqual(estimacion.normalizar_dedicacion(80), 0.8)
        self.assertEqual(estimacion.normalizar_dedicacion("20%"), 0.2)

    def test_rechaza_dedicacion_fuera_de_rango(self):
        for valor in (0, -10, 120, None):
            with self.subTest(valor=valor):
                with self.assertRaises(ValueError):
                    estimacion.normalizar_dedicacion(valor)

    def test_acepta_fecha_iso_y_dia_mes_anio(self):
        esperado = date(2026, 10, 5)
        self.assertEqual(estimacion._parse_fecha("2026-10-05"), esperado)
        self.assertEqual(estimacion._parse_fecha("05/10/2026"), esperado)

    def test_contexto_infiere_lider_y_responsable_desde_roles(self):
        fases = {
            "inception_fin_real": "2026-10-01",
            "construccion_inicio_real": "2026-10-05",
            "personas": [{
                "nombre": "Persona A",
                "dedicacion": "80%",
                "roles": [
                    "lider_proyecto", "responsable_documentacion", "desarrollador",
                ],
            }],
        }
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = Path(carpeta) / "estimacion-fases.json"
            ruta.write_text(json.dumps(fases), encoding="utf-8")
            contexto, modo = estimacion.resolver_contexto(
                state_path=Path(carpeta) / "sin-estado.md",
                fases_path=ruta,
            )
        self.assertEqual(modo, "completa")
        self.assertIsNone(contexto["lider"])
        estimacion.aplicar_equipo(contexto)
        self.assertEqual(estimacion.LIDER["name"], "Persona A")
        self.assertIs(estimacion.LIDER, estimacion.RESPONSABLE)


class CalendarioTest(unittest.TestCase):
    def test_venezuela_carga_feriados_fuera_del_rango_historico(self):
        self.assertFalse(estimacion.es_laborable(date(2032, 10, 12)))

    def test_otro_pais_exige_feriados_configurados(self):
        with self.assertRaises(ValueError):
            estimacion.validar_calendario("CR", set())
        etiqueta = estimacion.validar_calendario("CR", {date(2026, 1, 1)})
        self.assertIn("CR", etiqueta)


class AsignacionTest(unittest.TestCase):
    def _aplicar(self, *, exclusivo=False, modo="por_tarea", dependencias=None):
        estimacion.aplicar_equipo({
            "recursos": [
                {"nombre": "Persona B", "dedicacion": "100%", "agentes": 1},
                {"nombre": "Persona A", "dedicacion": 80, "agentes": 1},
                {
                    "nombre": "Persona C",
                    "dedicacion": "20%",
                    "agentes": 1,
                    "entregables": ["Integracion con el ERP"],
                    "exclusivo": exclusivo,
                },
            ],
            "lider": {"nombre": "Persona D", "dedicacion": 1},
            "responsable": {"nombre": "Persona D", "dedicacion": 1},
            "dependencias": dependencias or {},
            "modo_asignacion": modo,
        })

    def _programar(self, entregables):
        salida = []
        estimacion._programar_fase(
            entregables,
            estimacion.construir_dias_habiles(date(2026, 10, 5)),
            salida,
            {"EST": 0, "PAR": 0},
            0,
        )
        return salida

    def _aplicar_roles(self, *, lider_desarrolla=False, persona_c_inicio="2026-10-05"):
        personas = [
            {
                "nombre": "Persona D", "dedicacion": "20%",
                "roles": ["lider_proyecto", "responsable_documentacion"],
                "ejecuta_desarrollo": False,
            },
            {
                "nombre": "Persona A", "dedicacion": "80%",
                "roles": ["especialista_tecnico", "desarrollador"],
                "ejecuta_desarrollo": True,
            },
            {
                "nombre": "Persona B", "dedicacion": "100%",
                "roles": ["programador_principal", "desarrollador"],
                "ejecuta_desarrollo": True,
            },
            {
                "nombre": "Persona C", "dedicacion": "20%",
                "fecha_inicio_participacion": persona_c_inicio,
                "roles": ["especialista_integracion", "desarrollador"],
                "ejecuta_desarrollo": True,
                "asignacion": {
                    "tipo": "entregables",
                    "entregables": ["Integracion con el ERP"],
                    "exclusivo": True,
                },
            },
        ]
        if lider_desarrolla:
            personas = [p for p in personas if p["nombre"] != "Persona D"]
            personas[0]["roles"].extend(["lider_proyecto", "responsable_documentacion"])
        estimacion.aplicar_equipo({
            "recursos": personas,
            "modo_asignacion": "por_tarea",
            "roles_por_tarea": {
                "Definir especificaciones tecnicas": ["especialista_tecnico"],
                "Preparar ficha de la solicitud": ["responsable_documentacion"],
            },
            "responsables_funcionales": [
                {"nombre": "Usuario de negocio", "area": "Comercial"},
            ],
        })

    def test_por_tarea_distribuye_una_unidad_entre_persona_a_y_b(self):
        self._aplicar()
        salida = self._programar([
            ("Desarrollo completo", [
                ("Actividad A", 40, "media"),
                ("Actividad B", 40, "media"),
            ]),
        ])
        responsables = {t["recurso"] for t in salida[0]["tareas"]}
        self.assertEqual(responsables, {"Persona B", "Persona A"})

    def test_por_entregable_preserva_un_solo_dueno(self):
        self._aplicar(modo="por_entregable")
        salida = self._programar([
            ("Desarrollo completo", [
                ("Actividad A", 40, "media"),
                ("Actividad B", 40, "media"),
            ]),
        ])
        self.assertEqual(
            {t["recurso"] for t in salida[0]["tareas"]},
            {"Persona B"},
        )

    def test_especialista_no_exclusivo_participa_y_admite_apoyo(self):
        self._aplicar(exclusivo=False)
        salida = self._programar([
            ("Integracion con el ERP", [
                ("Cambio idzona", 40, "media"),
                ("Integracion de tablas nuevas", 40, "media"),
                ("Pruebas ERP", 40, "media"),
            ]),
        ])
        responsables = [t["recurso"] for t in salida[0]["tareas"]]
        self.assertEqual(responsables[0], "Persona C")
        self.assertIn("Persona B", responsables)
        self.assertIn("Persona A", responsables)

    def test_especialista_exclusivo_recibe_todo_su_entregable(self):
        self._aplicar(exclusivo=True)
        salida = self._programar([
            ("Integracion con el ERP", [
                ("Cambio idzona", 40, "media"),
                ("Integracion de tablas nuevas", 40, "media"),
            ]),
        ])
        self.assertEqual(
            {t["recurso"] for t in salida[0]["tareas"]},
            {"Persona C"},
        )

    def test_conserva_dependencias_entre_entregables(self):
        self._aplicar(dependencias={"Modulo B": ["Modulo A"]})
        salida = self._programar([
            ("Modulo A", [("Actividad A", 40, "media")]),
            ("Modulo B", [("Actividad B", 40, "media")]),
        ])
        self.assertGreaterEqual(salida[1]["inicio"], salida[0]["fin"])
        self.assertTrue(salida[1]["pred_jids"])

    def test_tres_frentes_paralelos_conservan_sus_candidatos(self):
        self._aplicar(exclusivo=False)
        salida = self._programar([
            ("Desarrollo A", [("Actividad A", 40, "media")]),
            ("Desarrollo B", [("Actividad B", 40, "media")]),
            ("Integracion con el ERP", [("Cambio idzona", 40, "media")]),
        ])
        por_nombre = {
            e["nombre"]: {t["recurso"] for t in e["tareas"]}
            for e in salida
        }
        self.assertNotIn("Persona C", por_nombre["Desarrollo A"])
        self.assertNotIn("Persona C", por_nombre["Desarrollo B"])
        self.assertIn("Persona C", por_nombre["Integracion con el ERP"])
        self.assertEqual({e["inicio"] for e in salida}, {date(2026, 10, 5)})

    def test_roles_unifican_persona_y_excluyen_lider_no_tecnico(self):
        self._aplicar_roles()
        self.assertIs(estimacion.LIDER, estimacion.RESPONSABLE)
        self.assertEqual(len(estimacion.RESOURCES), 4)
        self.assertNotIn(
            "Persona D",
            {p["name"] for p in estimacion._candidatos("Desarrollo general")},
        )
        self.assertEqual(len(estimacion.RESPONSABLES_FUNCIONALES), 1)
        self.assertNotIn(
            "Usuario de negocio", {r["name"] for r in estimacion.RESOURCES}
        )

    def test_hoja_colaboradores_documenta_roles_inicio_y_alcance(self):
        self._aplicar_roles()
        salida = self._programar([(
            "Integracion con el ERP",
            [("Construir integracion", 8, "media")],
        )])
        ws = estimacion.openpyxl.Workbook().active
        estimacion.write_colaboradores(ws, salida)
        encabezados = [ws.cell(row=3, column=i).value for i in range(1, 10)]
        self.assertIn("Roles", encabezados)
        self.assertIn("Inicio participación", encabezados)
        self.assertIn("Alcance", encabezados)
        nombres = [ws.cell(row=fila, column=2).value for fila in range(4, 8)]
        self.assertEqual(len(nombres), len(set(nombres)))

    def test_rol_requerido_asigna_definicion_al_especialista(self):
        self._aplicar_roles()
        salida = self._programar([(
            "Desarrollo general",
            [("Definir especificaciones tecnicas", 8, "discovery")],
        )])
        self.assertEqual(salida[0]["tareas"][0]["recurso"], "Persona A")

    def test_responsable_documentacion_recibe_tarea_sin_habilitarlo_como_desarrollador(self):
        self._aplicar_roles()
        salida = self._programar([(
            "Preparacion",
            [
                ("Preparar ficha de la solicitud", 4, "doc_funcional"),
                ("Construir cambio", 8, "media"),
            ],
        )])
        por_tarea = {t["nombre"]: t["recurso"] for t in salida[0]["tareas"]}
        self.assertEqual(por_tarea["Preparar ficha de la solicitud"], "Persona D")
        self.assertNotEqual(por_tarea["Construir cambio"], "Persona D")

    def test_fecha_de_inicio_limita_participacion(self):
        self._aplicar_roles(persona_c_inicio="2026-10-12")
        salida = self._programar([(
            "Integracion con el ERP",
            [("Construir integracion", 8, "media")],
        )])
        self.assertEqual(salida[0]["tareas"][0]["recurso"], "Persona C")
        self.assertGreaterEqual(salida[0]["inicio"], date(2026, 10, 12))

    def test_gestion_y_desarrollo_comparten_capacidad_si_es_la_misma_persona(self):
        self._aplicar_roles(lider_desarrolla=True)
        salida = self._programar([(
            "Desarrollo general",
            [
                ("Realizar el seguimiento y la gestion del entregable", 2, "gestion"),
                ("Definir especificaciones tecnicas", 8, "discovery"),
            ],
        )])
        gestion = next(t for t in salida[0]["tareas"] if t["tipo"] == "gestion")
        tarea_persona_a = next(
            t for t in salida[0]["tareas"]
            if t["tipo"] != "gestion" and t["recurso"] == "Persona A"
        )
        self.assertGreaterEqual(tarea_persona_a["h_ini"], gestion["h_fin"])

    def test_rechaza_dedicaciones_distintas_para_la_misma_persona(self):
        with self.assertRaises(ValueError):
            estimacion.aplicar_equipo({
                "recursos": [{
                    "nombre": "Persona A", "dedicacion": "80%",
                    "roles": ["lider_proyecto", "desarrollador"],
                }],
                "lider": {"nombre": "Persona A", "dedicacion": "100%"},
            })

    def test_rechaza_persona_duplicada_en_lugar_de_duplicar_capacidad(self):
        with self.assertRaises(ValueError):
            estimacion.aplicar_equipo({
                "recursos": [
                    {"nombre": "Persona A", "dedicacion": "80%", "roles": ["especialista_tecnico"]},
                    {"nombre": "Persona A", "dedicacion": "80%", "roles": ["desarrollador"]},
                ],
            })


class ConstruccionMenorTest(unittest.TestCase):
    """Fuente de construccion desde construccion-menor.json (cambios sin tracker)."""

    def _escribir(self, carpeta, datos):
        ruta = Path(carpeta) / "construccion-menor.json"
        ruta.write_text(json.dumps(datos), encoding="utf-8")
        return ruta

    def test_parse_lista_de_paquetes_conserva_orden_y_tareas(self):
        datos = {"paquetes": [
            {"entregable": "Cambio en el modulo A", "tareas": [
                "Implementar el cambio en el modulo A",
                "Verificacion tecnica de Cambio en el modulo A",
            ]},
            {"entregable": "Pruebas tecnicas y funcionales", "tareas": [
                "Escribir pruebas unitarias",
                "Registrar evidencias de las pruebas",
            ]},
        ]}
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = self._escribir(carpeta, datos)
            unidades = estimacion.parse_construccion_ficha(ruta)
        self.assertEqual([u["unidad"] for u in unidades],
                         ["Cambio en el modulo A", "Pruebas tecnicas y funcionales"])
        # La tarea "Verificacion tecnica..." se separa como fila de verificacion.
        self.assertEqual([t["nombre"] for t in unidades[0]["tareas"]],
                         ["Implementar el cambio en el modulo A"])
        self.assertIsNotNone(unidades[0]["verif"])
        self.assertEqual(unidades[0]["verif"]["id"], "")
        self.assertIsNone(unidades[1]["verif"])

    def test_parse_acepta_objeto_unico(self):
        datos = {"entregable": "Ajuste menor", "tareas": ["Implementar el ajuste"]}
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = self._escribir(carpeta, datos)
            unidades = estimacion.parse_construccion_ficha(ruta)
        self.assertEqual(len(unidades), 1)
        self.assertEqual(unidades[0]["unidad"], "Ajuste menor")

    def test_parse_paquete_sin_tareas_cae_en_minimo_util(self):
        datos = {"paquetes": [{"entregable": "Ajuste menor"}]}
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = self._escribir(carpeta, datos)
            unidades = estimacion.parse_construccion_ficha(ruta)
        nombres = [t["nombre"] for t in unidades[0]["tareas"]]
        self.assertIn("Implementar Ajuste menor", nombres)
        self.assertIsNotNone(unidades[0]["verif"])

    def test_parse_devuelve_vacio_si_no_existe(self):
        self.assertEqual(
            estimacion.parse_construccion_ficha(Path("no-existe-xyz.json")), [])

    def test_estructura_compatible_con_construir_entregables_construccion(self):
        datos = {"entregable": "Ajuste menor", "tareas": [
            "Implementar el ajuste", "Verificacion tecnica del cambio"]}
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = self._escribir(carpeta, datos)
            unidades = estimacion.parse_construccion_ficha(ruta)
        # Debe consumirse igual que la salida del tracker, sin lanzar.
        entregables = estimacion.construir_entregables_construccion(unidades)
        self.assertEqual(entregables[0][0], "Ajuste menor")
        self.assertTrue(entregables[0][1])  # tiene tareas
        self.assertIsNotNone(entregables[0][2])  # verif_meta presente

    def test_recorrido_corto_desde_bloque_dedicado(self):
        datos = {
            "entregable": "Ajuste menor", "tareas": ["Implementar el ajuste"],
            "recorrido_corto": {
                "objetivo": "Objetivo X.", "donde_vive": "Modulo Y",
                "criterio_termino": "Responde 200.",
            },
        }
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = self._escribir(carpeta, datos)
            recorrido = estimacion._recorrido_corto_ficha(ruta)
        self.assertEqual(recorrido["objetivo"], "Objetivo X.")
        self.assertEqual(recorrido["criterio_termino"], "Responde 200.")

    def test_meta_fuente_ficha_agrega_filas_de_trazabilidad(self):
        filas = estimacion._meta_fuente_construccion(
            "ficha", {"objetivo": "Obj.", "donde_vive": "Mod", "criterio_termino": "Crit."})
        etiquetas = [f[0] for f in filas]
        self.assertIn("Fuente de construcción", etiquetas)
        self.assertIn("Objetivo del cambio", etiquetas)
        self.assertIn("Criterio de término", etiquetas)

    def test_meta_fuente_vacia_solo_informa_fuente(self):
        filas = estimacion._meta_fuente_construccion("vacia", {})
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0][0], "Fuente de construcción")

    def test_precedencia_tracker_gana_sobre_ficha(self):
        tracker = (
            "## Unidad: infrastructure\n"
            "| ID | Tarea | Estado | Inicio | Fin |\n"
            "| --- | --- | --- | --- | --- |\n"
            "| INF-T01 | Setup del repositorio | \u2705 | | |\n"
        )
        ficha = {"entregable": "Ajuste menor", "tareas": ["Implementar el ajuste"]}
        with tempfile.TemporaryDirectory() as carpeta:
            ruta_tracker = Path(carpeta) / "INDICE-PROGRESO.md"
            ruta_tracker.write_text(tracker, encoding="utf-8")
            ruta_ficha = self._escribir(carpeta, ficha)
            orig_tracker, orig_ficha = estimacion.TRACKER_PATH, estimacion.CONSTRUCCION_MENOR_PATH
            try:
                estimacion.TRACKER_PATH = ruta_tracker
                estimacion.CONSTRUCCION_MENOR_PATH = ruta_ficha
                unidades, fuente, _ = estimacion.resolver_construccion()
            finally:
                estimacion.TRACKER_PATH = orig_tracker
                estimacion.CONSTRUCCION_MENOR_PATH = orig_ficha
        self.assertEqual(fuente, "tracker")
        self.assertEqual(unidades[0]["unidad"], "infrastructure")

    def test_precedencia_usa_ficha_cuando_no_hay_tracker(self):
        ficha = {"entregable": "Ajuste menor", "tareas": ["Implementar el ajuste"]}
        with tempfile.TemporaryDirectory() as carpeta:
            ruta_ficha = self._escribir(carpeta, ficha)
            orig_tracker, orig_ficha = estimacion.TRACKER_PATH, estimacion.CONSTRUCCION_MENOR_PATH
            try:
                estimacion.TRACKER_PATH = Path(carpeta) / "sin-tracker.md"
                estimacion.CONSTRUCCION_MENOR_PATH = ruta_ficha
                unidades, fuente, _ = estimacion.resolver_construccion()
            finally:
                estimacion.TRACKER_PATH = orig_tracker
                estimacion.CONSTRUCCION_MENOR_PATH = orig_ficha
        self.assertEqual(fuente, "ficha")
        self.assertEqual(unidades[0]["unidad"], "Ajuste menor")

    def test_precedencia_vacia_sin_tracker_ni_ficha(self):
        with tempfile.TemporaryDirectory() as carpeta:
            orig_tracker, orig_ficha = estimacion.TRACKER_PATH, estimacion.CONSTRUCCION_MENOR_PATH
            try:
                estimacion.TRACKER_PATH = Path(carpeta) / "sin-tracker.md"
                estimacion.CONSTRUCCION_MENOR_PATH = Path(carpeta) / "sin-ficha.json"
                unidades, fuente, _ = estimacion.resolver_construccion()
            finally:
                estimacion.TRACKER_PATH = orig_tracker
                estimacion.CONSTRUCCION_MENOR_PATH = orig_ficha
        self.assertEqual(fuente, "vacia")
        self.assertEqual(unidades, [])


class DependenciasExternasTest(unittest.TestCase):
    def test_normalizar_dependencia_externa_valida(self):
        dep = {
            "area": "Base de datos",
            "necesidad": "Aprovisionar PostgreSQL",
            "contacto": "DBA",
            "estado": "en gestion",
            "bloquea": "Modulo de persistencia",
        }
        res = estimacion.normalizar_dependencia_externa(dep)
        self.assertEqual(res["area"], "Base de datos")
        self.assertEqual(res["necesidad"], "Aprovisionar PostgreSQL")
        self.assertEqual(res["contacto"], "DBA")
        self.assertEqual(res["estado"], "en gestión")
        self.assertEqual(res["bloquea"], "Modulo de persistencia")

    def test_normalizar_dependencia_externa_valores_por_defecto(self):
        dep = {
            "area": "Infraestructura",
            "necesidad": "VPC y subredes",
        }
        res = estimacion.normalizar_dependencia_externa(dep)
        self.assertEqual(res["contacto"], "")
        self.assertEqual(res["estado"], "pendiente")
        self.assertIsNone(res["bloquea"])

    def test_normalizar_dependencia_externa_rechaza_invalida(self):
        for invalido in ({}, {"area": ""}, {"necesidad": ""}, "no-dict", None):
            with self.subTest(invalido=invalido):
                with self.assertRaises((ValueError, TypeError)):
                    estimacion.normalizar_dependencia_externa(invalido)

    def test_derivar_dependencias_desde_project_profile_rds_serverless(self):
        perfil = {
            "schemaVersion": 1,
            "infrastructure": ["rds", "serverless"],
        }
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = Path(carpeta) / "project-profile.json"
            ruta.write_text(json.dumps(perfil), encoding="utf-8")
            deps = estimacion.derivar_dependencias_desde_perfil(ruta)

        areas = [d["area"] for d in deps]
        self.assertIn("Base de datos", areas)
        self.assertIn("Infraestructura", areas)
        self.assertIn("AWS / Nube", areas)

    def test_derivar_dependencias_desde_project_profile_none(self):
        perfil = {
            "schemaVersion": 1,
            "infrastructure": ["none"],
        }
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = Path(carpeta) / "project-profile.json"
            ruta.write_text(json.dumps(perfil), encoding="utf-8")
            deps = estimacion.derivar_dependencias_desde_perfil(ruta)
        self.assertEqual(deps, [])

    def test_derivar_dependencias_desde_project_profile_undecided(self):
        perfil = {
            "schemaVersion": 1,
            "infrastructure": ["undecided"],
        }
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = Path(carpeta) / "project-profile.json"
            ruta.write_text(json.dumps(perfil), encoding="utf-8")
            deps = estimacion.derivar_dependencias_desde_perfil(ruta)
        # undecided se resuelve a una de las seis areas canonicas: Infraestructura.
        self.assertEqual(len(deps), 1)
        self.assertEqual(deps[0]["area"], "Infraestructura")

    def test_derivar_dependencias_desde_contenido_detecta_areas_por_wbs(self):
        unidades = [
            {"unidad": "Persistencia", "tareas": [
                {"nombre": "Crear esquema y tablas en la base de datos", "id": "T1"},
            ]},
            {"unidad": "Integraciones", "tareas": [
                {"nombre": "Configurar reglas de firewall y VPN entre sedes", "id": "T2"},
                {"nombre": "Cifrar secretos y credenciales sensibles", "id": "T3"},
                {"nombre": "Publicar función Lambda y bucket S3", "id": "T4"},
                {"nombre": "Integrar pasarela de pago del proveedor externo", "id": "T5"},
            ]},
        ]
        deps = estimacion.derivar_dependencias_desde_contenido(unidades)
        areas = {d["area"] for d in deps}
        self.assertIn("Base de datos", areas)
        self.assertIn("Comunicaciones / Redes", areas)
        self.assertIn("Seguridad", areas)
        self.assertIn("AWS / Nube", areas)
        self.assertIn("Proveedor / Terceros", areas)
        # Solo se recomiendan las seis areas canonicas, sin inventar otras.
        self.assertTrue(areas.issubset({d["area"] for d in estimacion.AREAS_SOPORTE.values()}))

    def test_derivar_dependencias_desde_contenido_sin_senales_no_agrega(self):
        unidades = [
            {"unidad": "Maquetado", "tareas": [
                {"nombre": "Ajustar estilos del formulario de contacto", "id": "T1"},
            ]},
        ]
        self.assertEqual(estimacion.derivar_dependencias_desde_contenido(unidades), [])
        self.assertEqual(estimacion.derivar_dependencias_desde_contenido([]), [])

    def test_fusionar_dependencias_respeta_precedencia_y_deduplica(self):
        declarada = {
            "area": "Base de datos",
            "necesidad": "Aprovisionar Postgres para el modulo X",
            "contacto": "DBA del proyecto",
            "estado": "en gestión",
            "bloquea": "Persistencia",
        }
        contenido = estimacion.derivar_dependencias_desde_contenido([
            {"unidad": "Persistencia", "tareas": [
                {"nombre": "Crear tablas en la base de datos", "id": "T1"},
                {"nombre": "Revisar secretos y cifrado", "id": "T2"},
            ]},
        ])
        fusion = estimacion.fusionar_dependencias([declarada], contenido)
        por_area = {estimacion._clave_area(d["area"]): d for d in fusion}
        # Base de datos declarada gana: conserva contacto/estado/bloquea propios.
        self.assertEqual(por_area["base_datos"]["contacto"], "DBA del proyecto")
        self.assertEqual(por_area["base_datos"]["estado"], "en gestión")
        self.assertEqual(por_area["base_datos"]["bloquea"], "Persistencia")
        # Seguridad, que solo vino por contenido, se agrega.
        self.assertIn("seguridad", por_area)
        # No se duplica el area de base de datos.
        claves = [estimacion._clave_area(d["area"]) for d in fusion]
        self.assertEqual(len(claves), len(set(claves)))

    def test_inyectar_tareas_enlace_en_habilitacion_entorno_mayor(self):
        entregables = [
            ("Mob Elaboration", [("Tarea 1", 4, "discovery")]),
            ("Habilitación del Entorno de Trabajo", [
                ("Habilitar acceso BD", 4, "despliegue"),
                ("Seguimiento y gestion", 1, "gestion"),
            ]),
            ("Implementación", [("Despliegue", 8, "implementacion")]),
        ]
        deps_ext = [
            {
                "area": "Base de datos",
                "necesidad": "Aprovisionar Postgres",
                "contacto": "DBA",
                "estado": "pendiente",
                "bloquea": None,
            }
        ]
        inyectados = estimacion.inyectar_tareas_enlace_soporte(entregables, deps_ext)
        entorno = next(e for e in inyectados if e[0] == "Habilitación del Entorno de Trabajo")
        nombres_tareas = [t[0] for t in entorno[1]]
        self.assertTrue(any("Contactar al área de Base de Datos" in n for n in nombres_tareas))
        self.assertIn("Contactar al área de Base de Datos", nombres_tareas[0])

    def test_inyectar_tareas_enlace_en_primer_entregable_menor(self):
        entregables = [
            ("Módulo de Facturación", [
                ("Crear endpoint", 8, "media"),
            ]),
            ("Implementación", [("Despliegue", 8, "implementacion")]),
        ]
        deps_ext = [
            {
                "area": "Infraestructura",
                "necesidad": "VPC y subredes",
                "contacto": "Infra",
                "estado": "pendiente",
                "bloquea": None,
            }
        ]
        inyectados = estimacion.inyectar_tareas_enlace_soporte(entregables, deps_ext)
        primer_ent = inyectados[0]
        nombres_tareas = [t[0] for t in primer_ent[1]]
        self.assertTrue(any("Contactar al área de Infraestructura" in n for n in nombres_tareas))
        self.assertIn("Contactar al área de Infraestructura", nombres_tareas[0])

    def test_bloqueo_agrega_predecesora_en_dependencias(self):
        entregables = [
            ("Habilitación del Entorno de Trabajo", [
                ("Contactar al área de Base de Datos y acordar el aprovisionamiento", 1.0, "enlace_soporte"),
            ]),
            ("Módulo de Persistencia", [
                ("Crear tablas", 8, "media"),
            ]),
        ]
        deps_ext = [
            {
                "area": "Base de datos",
                "necesidad": "Aprovisionar Postgres",
                "contacto": "DBA",
                "estado": "pendiente",
                "bloquea": "Módulo de Persistencia",
            }
        ]
        orig_deps = dict(estimacion.DEPENDENCIAS)
        try:
            estimacion.DEPENDENCIAS = {}
            estimacion.vincular_bloqueos_dependencias_externas(entregables, deps_ext)
            clave_destino = estimacion._norm("Módulo de Persistencia")
            self.assertIn(clave_destino, estimacion.DEPENDENCIAS)
            clave_origen = estimacion._norm("Habilitación del Entorno de Trabajo")
            self.assertIn(clave_origen, estimacion.DEPENDENCIAS[clave_destino])
        finally:
            estimacion.DEPENDENCIAS = orig_deps

    def test_write_resumen_renderiza_seccion_dependencias_externas(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Resumen"
        scheduled = [
            {
                "nombre": "Módulo A",
                "tareas": [
                    {
                        "nombre": "Tarea 1", "hh_gen": 0.1, "hh_sup": 0.8,
                        "hh_total": 0.9, "tipo": "media", "id_jerarquico": "PAR-001-001",
                        "inicio": date(2026, 10, 5), "fin": date(2026, 10, 6),
                    }
                ],
                "id_jerarquico": "PAR-001",
                "hh_gen": 0.1, "hh_sup": 0.8, "hh_total": 0.9, "hh_trad": 8.0,
                "latencia_horas": 0.0, "tiene_verificacion": False,
                "inicio": date(2026, 10, 5), "fin": date(2026, 10, 6),
            }
        ]
        deps_ext = [
            {
                "area": "Base de datos",
                "necesidad": "Aprovisionar PostgreSQL",
                "contacto": "DBA Pedro",
                "estado": "pendiente",
                "bloquea": "Módulo A",
            }
        ]
        estimacion.write_resumen(ws, scheduled, dependencias_externas=deps_ext)
        texto_celdas = [str(ws.cell(row=r, column=c).value or "") for r in range(1, ws.max_row + 1) for c in range(1, 7)]
        self.assertTrue(any("Dependencias Externas de Áreas de Soporte" in t for t in texto_celdas))
        self.assertTrue(any("Base de datos" in t for t in texto_celdas))
        self.assertTrue(any("Aprovisionar PostgreSQL" in t for t in texto_celdas))
        self.assertTrue(any("DBA Pedro" in t for t in texto_celdas))


if __name__ == "__main__":
    unittest.main()
